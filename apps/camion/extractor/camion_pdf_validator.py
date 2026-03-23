#!/usr/bin/env python3
from __future__ import annotations

import argparse
import dataclasses
import datetime as dt
import re
import unicodedata
from pathlib import Path
from typing import Iterable, Optional

import fitz
import openpyxl
import pytesseract
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image
from pypdf import PdfReader

MISSING_FILL = PatternFill(fill_type="solid", fgColor="FFF2F2")
CONFIRMED_FILL = PatternFill(fill_type="solid", fgColor="F2FFF2")

_COL = {
    'tour': 0,
    'date': 1,
    'trailer': 2,
    'to': 3,
    'shipper_name': 4,
    'shipper_iso': 5,
    'shipper_city': 6,
    'shipper_code': 7,
    'recipient_name': 8,
    'recipient_iso': 9,
    'recipient_city': 10,
    'recipient_code': 11,
    'vol': 12,
    'mrn_invoice': 13,
    'mrn_detail': 14,
    'value_eur': 15,
    'value_usd': 16,
    'hu': 17,
    'peso_bruto': 18,
    'peso_neto': 19,
    'pk': 20,
    'cl': 21,
}


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace("\n", " ")
    text = re.sub(r"\s+", " ", text).strip().lower()
    return text


def normalize_token(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", normalize_text(value))


def normalize_words(value: object) -> list[str]:
    text = normalize_text(value)
    words = re.findall(r"[a-z0-9]+", text)
    stop = {
        'the', 'and', 'for', 'with', 'from', 'city', 'corp', 'corporation', 'automotive',
        'srl', 'sas', 'sl', 'gmbh', 'soc', 'unipers', 'make', 'route', 'nationale'
    }
    return [w for w in words if len(w) >= 3 and w not in stop]


def float_candidates(value: Optional[float]) -> list[str]:
    if value is None:
        return []
    v = float(value)
    candidates = {
        str(int(round(v))) if v.is_integer() else f"{v:.2f}",
        f"{v:.2f}",
        f"{v:.1f}",
        f"{v:.3f}",
    }
    expanded = set()
    for item in candidates:
        expanded.add(item)
        expanded.add(item.replace('.', ','))
        expanded.add(item.replace(',', '.'))
        expanded.add(item.replace('.', ''))
        expanded.add(item.replace(',', ''))
    return sorted(x for x in expanded if x)


def safe_float(value: object) -> Optional[float]:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except Exception:
        return None



def safe_get(row, idx, default=None):
    if row is None:
        return default
    if idx < 0 or idx >= len(row):
        return default
    return row[idx]

def _find_source_worksheet(wb: openpyxl.Workbook):
    for ws in wb.worksheets:
        try:
            row1 = [str(ws.cell(1, c).value or "").strip().lower() for c in range(1, 13)]
            has_base = (
                len(row1) >= 4
                and row1[0] == 'tour'
                and row1[1] == 'date'
                and row1[2] == 'trailer'
                and row1[3] == 'to'
            )
            has_shipper = 'shipper' in row1
            has_recipient = 'recipient' in row1
            has_some_data = any(
                any(ws.cell(r, c).value not in (None, '') for c in range(1, 8))
                for r in range(5, min(ws.max_row, 15) + 1)
            )
            if has_base and has_shipper and has_recipient and has_some_data:
                return ws
        except Exception:
            continue
    raise ValueError('No se encontró ninguna hoja con formato válido de packing list')


@dataclasses.dataclass
class RowEntry:
    excel_row: int
    tour: Optional[str]
    date: Optional[str]
    trailer: Optional[str]
    to_code: Optional[str]
    shipper_name: Optional[str]
    shipper_iso: Optional[str]
    shipper_city: Optional[str]
    shipper_code: Optional[str]
    recipient_name: Optional[str]
    recipient_iso: Optional[str]
    recipient_city: Optional[str]
    recipient_code: Optional[str]
    vol: Optional[float]
    mrn_invoice: Optional[str]
    mrn_detail: Optional[str]
    value_eur: Optional[float]
    value_usd: Optional[float]
    hu: Optional[float]
    peso_bruto: Optional[float]
    peso_neto: Optional[float]
    pk: Optional[float]
    cl: Optional[float]

    def split_tokens(self, raw: Optional[str]) -> list[str]:
        tokens: list[str] = []
        if not raw:
            return tokens
        for part in re.split(r"[\n,/;]+", str(raw)):
            part = part.strip()
            if not part:
                continue
            part = re.sub(r"^(EX|T1)\s*:\s*", "", part, flags=re.I)
            tok = normalize_token(part)
            if tok:
                tokens.append(tok)
        return sorted(set(tokens))

    @property
    def mrn_invoice_tokens(self) -> list[str]:
        return self.split_tokens(self.mrn_invoice)

    @property
    def mrn_detail_tokens(self) -> list[str]:
        return self.split_tokens(self.mrn_detail)

    @property
    def shipper_words(self) -> list[str]:
        return normalize_words(self.shipper_name) + normalize_words(self.shipper_city) + normalize_words(self.shipper_code)

    @property
    def recipient_words(self) -> list[str]:
        return normalize_words(self.recipient_name) + normalize_words(self.recipient_city) + normalize_words(self.recipient_code)


@dataclasses.dataclass
class PageMatch:
    page: int
    score: int
    reasons: list[str]
    matched_tokens: dict[str, object]


def load_entries_from_xlsx(xlsx_path: Path) -> list[RowEntry]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    ws = _find_source_worksheet(wb)
    entries: list[RowEntry] = []
    current_tour = current_date = current_trailer = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
        if not any(cell is not None and cell != "" for cell in row):
            continue

        tour, date_val, trailer = safe_get(row, 0), safe_get(row, 1), safe_get(row, 2)
        if tour:
            current_tour = str(tour)
        if date_val:
            if isinstance(date_val, (dt.datetime, dt.date)):
                current_date = date_val.strftime('%Y-%m-%d')
            else:
                current_date = str(date_val)
        if trailer:
            current_trailer = str(trailer)

        entries.append(RowEntry(
            excel_row=row_idx,
            tour=current_tour,
            date=current_date,
            trailer=current_trailer,
            to_code=str(safe_get(row, 3)).strip() if safe_get(row, 3) else None,
            shipper_name=str(safe_get(row, 4)).strip() if safe_get(row, 4) else None,
            shipper_iso=str(safe_get(row, 5)).strip() if safe_get(row, 5) else None,
            shipper_city=str(safe_get(row, 6)).strip() if safe_get(row, 6) else None,
            shipper_code=str(safe_get(row, 7)).strip() if safe_get(row, 7) else None,
            recipient_name=str(safe_get(row, 8)).strip() if safe_get(row, 8) else None,
            recipient_iso=str(safe_get(row, 9)).strip() if safe_get(row, 9) else None,
            recipient_city=str(safe_get(row, 10)).strip() if safe_get(row, 10) else None,
            recipient_code=str(safe_get(row, 11)).strip() if safe_get(row, 11) else None,
            vol=safe_float(safe_get(row, 12)),
            mrn_invoice=str(safe_get(row, 13)).strip() if safe_get(row, 13) else None,
            mrn_detail=str(safe_get(row, 14)).strip() if safe_get(row, 14) else None,
            value_eur=safe_float(safe_get(row, 15)),
            value_usd=safe_float(safe_get(row, 16)),
            hu=safe_float(safe_get(row, 17)),
            peso_bruto=safe_float(safe_get(row, 18)),
            peso_neto=safe_float(safe_get(row, 19)),
            pk=safe_float(safe_get(row, 20)),
            cl=safe_float(safe_get(row, 21)),
        ))
    return entries


def extract_pdf_text(pdf_path: Path, dpi: int = 150, ocr_lang: str = 'eng+spa') -> dict[int, str]:
    page_texts: dict[int, str] = {}
    native_texts: dict[int, str] = {}
    reader = PdfReader(str(pdf_path))
    for i, page in enumerate(reader.pages, start=1):
        try:
            native_texts[i] = page.extract_text() or ''
        except Exception:
            native_texts[i] = ''

    doc = fitz.open(str(pdf_path))
    for page_no in range(1, len(doc) + 1):
        native = native_texts.get(page_no, '')
        if len(normalize_text(native)) >= 80:
            page_texts[page_no] = native
            continue
        page = doc[page_no - 1]
        pix = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False)
        img = Image.frombytes('RGB', [pix.width, pix.height], pix.samples)
        text = pytesseract.image_to_string(img, lang=ocr_lang) or ''
        page_texts[page_no] = text
    return page_texts


def _find_present_values(haystack: str, needles: Iterable[str]) -> list[str]:
    found = []
    for n in needles:
        if n and n in haystack:
            found.append(n)
    return found


def _word_score(page_words: set[str], target_words: list[str]) -> tuple[int, list[str]]:
    target = sorted(set(w for w in target_words if len(w) >= 3))
    hits = [w for w in target if w in page_words]
    return len(hits), hits


def score_entry_on_page(entry: RowEntry, page_text: str) -> PageMatch | None:
    tok = normalize_token(page_text)
    words = set(normalize_words(page_text))
    score = 0
    reasons: list[str] = []
    matched: dict[str, object] = {}

    to_token = normalize_token(entry.to_code)
    if to_token and to_token in tok:
        score += 60
        reasons.append('TO')
        matched['to'] = True

    inv_hits = _find_present_values(tok, entry.mrn_invoice_tokens)
    if inv_hits:
        score += 30 + 10 * min(len(inv_hits), 3)
        reasons.append('invoice')
        matched['mrn_invoice'] = inv_hits

    det_hits = _find_present_values(tok, entry.mrn_detail_tokens)
    if det_hits:
        score += 24 + 8 * min(len(det_hits), 3)
        reasons.append('detail')
        matched['mrn_detail'] = det_hits

    shipper_count, shipper_hits = _word_score(words, entry.shipper_words)
    if shipper_count:
        score += min(shipper_count, 4) * 6
        reasons.append('shipper')
        matched['shipper_words'] = shipper_hits

    recipient_count, recipient_hits = _word_score(words, entry.recipient_words)
    if recipient_count:
        score += min(recipient_count, 4) * 6
        reasons.append('recipient')
        matched['recipient_words'] = recipient_hits

    for key, value, pts in [
        ('gross', entry.peso_bruto, 8),
        ('net', entry.peso_neto, 8),
        ('hu', entry.hu, 4),
        ('vol', entry.vol, 5),
        ('pk', entry.pk, 4),
        ('cl', entry.cl, 4),
        ('value_eur', entry.value_eur, 6),
        ('value_usd', entry.value_usd, 6),
    ]:
        hits = _find_present_values(tok, [normalize_token(x) for x in float_candidates(value)])
        if hits:
            score += pts
            reasons.append(key)
            matched[key] = hits[:3]

    if score == 0:
        return None
    return PageMatch(page=-1, score=score, reasons=reasons, matched_tokens=matched)


def find_best_pages(entry: RowEntry, page_texts: dict[int, str], top_n: int = 6) -> list[PageMatch]:
    matches = []
    for page_no, text in page_texts.items():
        pm = score_entry_on_page(entry, text)
        if pm is None:
            continue
        pm.page = page_no
        matches.append(pm)
    matches.sort(key=lambda x: (-x.score, x.page))
    return matches[:top_n]


def gather_context_text(matches: list[PageMatch], page_texts: dict[int, str], extra_window: int = 1) -> tuple[str, set[str], list[int]]:
    if not matches:
        return '', set(), []
    page_set = set()
    max_page = max(page_texts) if page_texts else 0
    for m in matches[:3]:
        for p in range(max(1, m.page - extra_window), min(max_page, m.page + extra_window) + 1):
            page_set.add(p)
    pages = sorted(page_set)
    text = '\n\n'.join(page_texts[p] for p in pages)
    return normalize_token(text), set(normalize_words(text)), pages


def token_present(ctx_token: str, value: object) -> bool:
    tok = normalize_token(value)
    return bool(tok) and tok in ctx_token


def fuzzy_words_present(ctx_words: set[str], value: object, min_hits: int = 1) -> bool:
    words = [w for w in normalize_words(value) if len(w) >= 3]
    if not words:
        return False
    hits = [w for w in words if w in ctx_words]
    need = min(min_hits, len(set(words)))
    if len(hits) >= need:
        return True
    if len(set(words)) >= 2 and len(hits) >= 2:
        return True
    return False


def any_float_present(ctx_token: str, value: Optional[float]) -> bool:
    if value is None:
        return False
    candidates = [normalize_token(x) for x in float_candidates(value)]
    return any(c and c in ctx_token for c in candidates)


def field_confirmation(entry: RowEntry, matches: list[PageMatch], page_texts: dict[int, str]) -> dict[str, object]:
    ctx_token, ctx_words, ctx_pages = gather_context_text(matches, page_texts, extra_window=1)
    return {
        'context_pages': ctx_pages,
        'A': False,
        'B': False,
        'C': False,
        'D': token_present(ctx_token, entry.to_code),
        'E': fuzzy_words_present(ctx_words, entry.shipper_name, min_hits=1),
        'F': token_present(ctx_token, entry.shipper_iso),
        'G': fuzzy_words_present(ctx_words, entry.shipper_city, min_hits=1),
        'H': token_present(ctx_token, entry.shipper_code),
        'I': fuzzy_words_present(ctx_words, entry.recipient_name, min_hits=1),
        'J': token_present(ctx_token, entry.recipient_iso),
        'K': fuzzy_words_present(ctx_words, entry.recipient_city, min_hits=1),
        'L': token_present(ctx_token, entry.recipient_code),
        'M': any_float_present(ctx_token, entry.vol),
        'N': any(token_present(ctx_token, t) for t in entry.mrn_invoice_tokens),
        'O': any(token_present(ctx_token, t) for t in entry.mrn_detail_tokens),
        'P': any_float_present(ctx_token, entry.value_eur),
        'Q': any_float_present(ctx_token, entry.value_usd),
        'R': any_float_present(ctx_token, entry.hu),
        'S': any_float_present(ctx_token, entry.peso_bruto),
        'T': any_float_present(ctx_token, entry.peso_neto),
        'U': any_float_present(ctx_token, entry.pk),
        'V': any_float_present(ctx_token, entry.cl),
    }


def infer_status(entry: RowEntry, matches: list[PageMatch], confirmations: dict[str, object]) -> tuple[str, list[str]]:
    if not matches:
        return 'NO_MATCH', ['No se localizaron páginas candidatas']
    best = matches[0]
    issues = []
    if best.score < 35:
        issues.append(f'Score bajo: {best.score}')
    if entry.to_code and not confirmations.get('D'):
        issues.append('No aparece el TO en el contexto')
    if entry.mrn_invoice and not confirmations.get('N'):
        issues.append('No aparece el MRN/Invoice')
    return ('REVIEW', issues) if issues else ('OK', [])


def build_report(entries: list[RowEntry], page_texts: dict[int, str]) -> dict:
    results = []
    summary = {'entries': len(entries), 'ok': 0, 'review': 0, 'no_match': 0, 'pdf_pages': len(page_texts)}
    for entry in entries:
        matches = find_best_pages(entry, page_texts)
        confirmations = field_confirmation(entry, matches, page_texts) if matches else {}
        status, issues = infer_status(entry, matches, confirmations)
        summary[status.lower()] += 1
        results.append({
            'excel_row': entry.excel_row,
            'status': status,
            'issues': issues,
            'best_pages': [m.page for m in matches],
            'confirmations': confirmations,
        })
    return {'summary': summary, 'results': results}


def add_validated_sheet(wb: openpyxl.Workbook, report: dict, validated_sheet_name: str = 'PDF_VALIDADO') -> None:
    src_ws = _find_source_worksheet(wb)
    if validated_sheet_name in wb.sheetnames:
        del wb[validated_sheet_name]
    dst_ws = wb.copy_worksheet(src_ws)
    dst_ws.title = validated_sheet_name

    row_map = {item['excel_row']: item for item in report['results']}
    review_col = 23
    pages_col = 24

    dst_ws.cell(1, review_col).value = 'Validation Status'
    dst_ws.cell(1, pages_col).value = 'PDF Pages'

    for row_idx in range(5, dst_ws.max_row + 1):
        item = row_map.get(row_idx)
        if not item:
            continue
        conf = item.get('confirmations', {})
        for col_idx in range(1, 23):
            col_letter = get_column_letter(col_idx)
            cell = dst_ws.cell(row_idx, col_idx)
            original_has_value = cell.value not in (None, '')
            keep = bool(conf.get(col_letter, False))
            if original_has_value and not keep:
                cell.value = None
                cell.fill = MISSING_FILL
            elif original_has_value and keep:
                cell.fill = CONFIRMED_FILL
        dst_ws.cell(row_idx, review_col).value = item.get('status')
        dst_ws.cell(row_idx, pages_col).value = ', '.join(map(str, conf.get('context_pages', [])))


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description='Validador XLSX vs DOC PDF para Camión Export')
    p.add_argument('--xlsx', required=True)
    p.add_argument('--doc', required=True)
    p.add_argument('--output', default=None)
    p.add_argument('--dpi', type=int, default=150)
    return p.parse_args()


def derive_output_path(xlsx_path: str, output: Optional[str]) -> Path:
    if output:
        return Path(output)
    src = Path(xlsx_path)
    return src.with_name(f'{src.stem}-VALIDATED.xlsx')

def validate_xlsx_against_doc(xlsx_path: Path, doc_path: Path, dpi: int = 150) -> dict:
    entries = load_entries_from_xlsx(Path(xlsx_path))
    page_texts = extract_pdf_text(Path(doc_path), dpi=dpi)
    return build_report(entries, page_texts)


def append_validated_sheet_from_paths(
    wb: openpyxl.Workbook,
    xlsx_path: Path,
    doc_path: Path,
    dpi: int = 150,
    validated_sheet_name: str = "PDF_VALIDADO",
) -> dict:
    report = validate_xlsx_against_doc(
        xlsx_path=Path(xlsx_path),
        doc_path=Path(doc_path),
        dpi=dpi,
    )
    add_validated_sheet(wb, report, validated_sheet_name=validated_sheet_name)
    return report

def main() -> int:
    args = parse_args()
    xlsx_path = Path(args.xlsx)
    doc_path = Path(args.doc)
    out_path = derive_output_path(args.xlsx, args.output)
    entries = load_entries_from_xlsx(xlsx_path)
    page_texts = extract_pdf_text(doc_path, dpi=args.dpi)
    report = build_report(entries, page_texts)
    wb = openpyxl.load_workbook(xlsx_path)
    add_validated_sheet(wb, report)
    wb.save(out_path)
    print(f'OK: {out_path}')
    print(report['summary'])
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
