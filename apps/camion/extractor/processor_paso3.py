#!/usr/bin/env python3
"""
processor_paso3.py  –  Hoja PASO_3 (DAE-only) y CSV PASO_4
==============================================================================
Toma las filas ya leídas de la fuente (vía processor_paso1.read_source_paso1)
y produce:

  1. Hoja PASO_3 en el XLSX:
        4 columnas: MRN (sin "EX:") | HS code | Bultos unif. | Peso Bruto
        - Solo filas cuyo MRN/Invoice empieza por "EX:".
        - Bultos unificados = SUMA de PK + BX + PX + CL + RO.
        - Peso Bruto redondeado half-up.
        - HS code rellenado desde el dict {MRN: HS} pasado como argumento;
          cadena vacía si no hay match.
        - Alternancia estricta amarillo/blanco.

  2. CSV PASO_4 (fichero independiente al lado del XLSX):
        Mismo contenido que PASO_3 pero:
        - Encoding cp1252 (Windows-1252 / WinLatin 1).
        - Separador ';'.
        - Line ending CRLF.
        - Sin comillas envolviendo campos (QUOTE_NONE).

Uso integrado:
    from processor_paso3 import build_paso3_rows, add_paso3_sheet, write_paso4_csv

    rows, _ = read_source_paso1(xlsx)
    paso3_rows = build_paso3_rows(rows, hs_map={'26FR10002980612MB3': '3926', ...})
    add_paso3_sheet(wb, paso3_rows)
    write_paso4_csv('/path/to/PL-PASO4.csv', paso3_rows)
"""

from __future__ import annotations

import csv
import math
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Font, PatternFill

# ─────────────────────────────────────────────────────────────────────────────
# Estilos (mismos códigos que processor_paso1)
# ─────────────────────────────────────────────────────────────────────────────

_YELLOW_DATA = 'FFFFFFBB'


def _fill(hex_rgb: str) -> PatternFill:
    return PatternFill('solid', start_color=hex_rgb, end_color=hex_rgb)


def _round_half_up(x: float) -> int:
    return math.floor(x + 0.5)


# ─────────────────────────────────────────────────────────────────────────────
# Filtro DAE + transformación
# ─────────────────────────────────────────────────────────────────────────────

def _is_dae_row(row: dict) -> bool:
    mrn = str(row.get('mrn_invoice') or '').strip()
    return mrn.upper().startswith('EX:')


def _strip_ex_prefix(mrn: str) -> str:
    """'EX: 26FR10002980612MB3' → '26FR10002980612MB3'"""
    s = str(mrn or '').strip()
    if s.upper().startswith('EX:'):
        s = s[3:].strip()
    return s


def _unified_bultos(row: dict) -> Optional[int]:
    """
    Suma todos los tipos de bulto de la fila (PK + BX + PX + CL + RO).

    Una fila DAE puede tener bultos repartidos en múltiples columnas
    (ej. 1 BX + 5 PX = 6 unidades totales). Anteriormente solo se leía
    la primera columna no-vacía, lo cual fallaba silenciosamente cuando
    había más de un tipo de bulto en la misma línea.

    Returns:
        Suma total de bultos como int, o None si todas las columnas están vacías.
    """
    total = 0
    found = False
    for key in ('pk', 'bx', 'px', 'cl', 'ro'):
        v = row.get(key)
        if v is None or v == '':
            continue
        try:
            total += int(v)
            found = True
        except (TypeError, ValueError):
            try:
                total += int(float(v))
                found = True
            except (TypeError, ValueError):
                pass
    return total if found else None


def build_paso3_rows(source_rows: list[dict],
                     hs_map: Optional[dict] = None) -> list[dict]:
    """
    Filtra las filas DAE y devuelve las 4 columnas listas para escribir.

    Args:
        source_rows: filas devueltas por read_source_paso1()
        hs_map:      Resultado de hs_extractor.extract_hs_map(). Acepta dos
                     formas para compatibilidad:
                       - Forma nueva: {'mrn': {MRN: HS4}, 'shipper': {key: HS4}}
                       - Forma antigua (v1): {MRN: HS4}
                     Cuando un MRN no está en mrn-map, intenta matchear por
                     nombre de shipper (caso EADs italianos donde el MRN del
                     código de barras sale mangled del OCR).

    Returns:
        Lista de dicts {mrn, hs, bultos, peso_bruto}
    """
    hs_map = hs_map or {}

    # Normalizar a forma nueva
    if 'mrn' in hs_map and 'shipper' in hs_map:
        mrn_map     = hs_map.get('mrn',     {}) or {}
        shipper_map = hs_map.get('shipper', {}) or {}
    else:
        # Forma antigua → tratar todo como mrn_map
        mrn_map     = hs_map
        shipper_map = {}

    out: list[dict] = []

    for row in source_rows:
        if not _is_dae_row(row):
            continue
        mrn = _strip_ex_prefix(row['mrn_invoice'])
        pb  = row['peso_bruto']

        # 1. Intentar match directo por MRN
        hs = mrn_map.get(mrn, '')

        # 2. Fallback: match por shipper si tenemos su nombre y un huérfano
        if not hs and shipper_map:
            shipper_name = (row.get('shipper_name') or '').upper()
            for ship_key, ship_hs in shipper_map.items():
                # Comparación robusta — keyword del extractor debe estar en el
                # shipper del XLSX (ej. "MASTROTTO" in "Gruppo Mastrotto SpA")
                if ship_key.upper() in shipper_name:
                    hs = ship_hs
                    break
                # Variantes con acentos / sufijos comunes
                normalized = (shipper_name
                              .replace('Ä', 'A').replace('Ü', 'U').replace('Ö', 'O'))
                if ship_key.upper() in normalized:
                    hs = ship_hs
                    break

        out.append({
            'mrn':        mrn,
            'hs':         hs,
            'bultos':     _unified_bultos(row),
            'peso_bruto': _round_half_up(pb) if pb is not None else None,
        })

    return out


# ─────────────────────────────────────────────────────────────────────────────
# Sheet writer
# ─────────────────────────────────────────────────────────────────────────────

def add_paso3_sheet(wb: openpyxl.Workbook,
                    rows: list[dict],
                    sheet_name: str = 'PASO_3') -> None:
    """
    Escribe la hoja PASO_3 con alternancia estricta amarillo/blanco.
    No tiene cabeceras ni totales (replica el formato manual del cliente).
    """
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    arial10 = Font(name='Arial', size=10)

    for i, row in enumerate(rows):
        ws.append([
            row['mrn'],
            row['hs'] or None,
            row['bultos'],
            row['peso_bruto'],
        ])
        excel_row_idx = ws.max_row     # ← capturar TRAS el append (max_row puede mentir antes)
        is_yellow = (i % 2) == 0   # R1 amarillo, R2 blanco — replica el manual
        for cell in ws[excel_row_idx]:
            cell.font = arial10
            if is_yellow:
                cell.fill = _fill(_YELLOW_DATA)

        # Formatos de celda como en el PASO_3 manual
        ws.cell(excel_row_idx, 1).number_format = '@'   # MRN como texto
        ws.cell(excel_row_idx, 2).number_format = '@'   # HS como texto (preserva ceros)
        ws.cell(excel_row_idx, 3).number_format = '0'   # bultos entero
        # C4 (peso bruto) queda en General

    # Anchos
    for col_letter, w in zip('ABCD', [22, 8, 8, 12]):
        ws.column_dimensions[col_letter].width = w


# ─────────────────────────────────────────────────────────────────────────────
# CSV writer (PASO_4)
# ─────────────────────────────────────────────────────────────────────────────

def write_paso4_csv(csv_path: str | Path, rows: list[dict]) -> Path:
    """
    Escribe el CSV PASO_4 con las características que pide "visual":
      - Encoding cp1252 (Windows-1252 / WinLatin 1)
      - Separador ';'
      - Line ending CRLF (\\r\\n)
      - Sin comillas (QUOTE_NONE)
      - Sin cabeceras

    Args:
        csv_path: ruta destino del CSV.
        rows:     mismas filas que devuelve build_paso3_rows().

    Returns:
        Path al fichero escrito.
    """
    csv_path = Path(csv_path)
    csv_path.parent.mkdir(parents=True, exist_ok=True)

    # newline='' es OBLIGATORIO en csv: deja que el writer ponga el lineterminator
    # correcto sin que el sistema lo duplique a \r\r\n en Windows.
    with open(csv_path, 'w', encoding='cp1252', errors='replace', newline='') as f:
        writer = csv.writer(
            f,
            delimiter=';',
            lineterminator='\r\n',
            quoting=csv.QUOTE_NONE,
            escapechar='\\',
        )
        for row in rows:
            writer.writerow([
                row['mrn'] or '',
                row['hs']  or '',
                '' if row['bultos']     is None else row['bultos'],
                '' if row['peso_bruto'] is None else row['peso_bruto'],
            ])

    return csv_path
