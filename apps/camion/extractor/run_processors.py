#!/usr/bin/env python3
"""
run_processors.py  –  Orquestador: pipelines DAE + T1 + validación
====================================================================
Ejecuta todos los pipelines sobre el mismo XLSX fuente y genera:

  Hojas en el XLSX de salida:
    PL_ORIGINAL          – hoja original del cliente (renombrada de Sheet1)
    VALIDACION_PL_DOC    – validación XLSX vs DOC.pdf (si se pasa --validate)
    DAE_CIERRE_CAMION    – cierre del camión (DAE pipeline): 9 cols,
                           todas las filas, PB redondeado half-up,
                           alternancia estricta, COINCIDE final
    DAE_LISTADO          – DAE-only: 4 cols (MRN sin EX:, HS, bultos, PB)
    T1_CIERRE_CAMION     – cierre del camión (T1 pipeline): mismas filas
                           pero con peso bruto del PDF T1 (reagrupa por MRN)
    T1_RESUMEN           – metadata extraída de los PDFs T1

  Fichero CSV junto al XLSX:
    <stem>-DAE_VISUAL.csv  – contenido de DAE_LISTADO en cp1252; CRLF;
                             sin quotes; para subir a Visual

Uso:
    python run_processors.py --xlsx PL.xlsx \\
                             --t1 T1_1.pdf T1_2.pdf \\
                             --doc DOC1.pdf DOC2.pdf \\
                             --validate \\
                             -o OUT/
"""

from __future__ import annotations

SCRIPT_VERSION = "2026-04-25.v6"

SCRIPT_CHANGELOG = """
## 2026-04-25.v6

### Renombrado de hojas (rompe compatibilidad con scripts externos)
- Sheet1                 → PL_ORIGINAL
- PASO_1                 → DAE_CIERRE_CAMION
- PASO_3                 → DAE_LISTADO
- T1_RESULT              → T1_CIERRE_CAMION
- T1_SUMMARY             → T1_RESUMEN
- PDF_VALIDADO           → VALIDACION_PL_DOC
- <stem>-PASO4.csv       → <stem>-DAE_VISUAL.csv

## 2026-04-23.v2

### Nuevo
- Hojas DAE_CIERRE_CAMION y DAE_LISTADO con layout del cliente.
- CSV DAE_VISUAL al lado del XLSX (cp1252, ;, CRLF, sin quotes).
- Extracción automática de HS codes desde DOC.pdf vía OCR (tesseract).

### Fixes
- camion_export_processor.py: detección dinámica de columnas BX/PX/CL/RO
  (formato Tangier-25-cols ya no lee BX como CL).
"""

import argparse
import sys
from pathlib import Path
from typing import Optional

import openpyxl

# ── Importar todos los procesadores ──────────────────────────────────────────
try:
    import camion_export_processor as t1_processor
except ImportError:
    print('\n❌ No se encuentra camion_export_processor.py en el mismo directorio.\n')
    sys.exit(1)

try:
    from processor_paso1 import (
        read_source_paso1, process_paso1, add_paso1_sheet,
    )
except ImportError:
    print('\n❌ No se encuentra processor_paso1.py en el mismo directorio.\n')
    sys.exit(1)

try:
    from processor_paso3 import (
        build_paso3_rows, add_paso3_sheet, write_paso4_csv,
    )
except ImportError:
    print('\n❌ No se encuentra processor_paso3.py en el mismo directorio.\n')
    sys.exit(1)


# ── Output path helpers ──────────────────────────────────────────────────────

def derive_output_xlsx(xlsx_path: str, output: Optional[str]) -> Path:
    src = Path(xlsx_path)

    if not output:
        return src.with_name(f'{src.stem}-PROCESSED.xlsx')

    out = Path(output)
    if out.suffix.lower() == '.xlsx':
        out.parent.mkdir(parents=True, exist_ok=True)
        return out
    if out.exists() and out.is_dir():
        return out / f'{src.stem}-PROCESSED.xlsx'
    if out.suffix == '':
        out.mkdir(parents=True, exist_ok=True)
        return out / f'{src.stem}-PROCESSED.xlsx'
    out.parent.mkdir(parents=True, exist_ok=True)
    return out


def derive_csv_path(xlsx_output: Path) -> Path:
    """PL-PROCESSED.xlsx → PL-DAE_VISUAL.csv (mismo directorio, mismo stem-base)."""
    stem = xlsx_output.stem
    if stem.endswith('-PROCESSED'):
        stem = stem[:-len('-PROCESSED')]
    return xlsx_output.with_name(f'{stem}-DAE_VISUAL.csv')


# ── Main logic ───────────────────────────────────────────────────────────────

def run(
    xlsx_path:       str,
    t1_paths:        list[str],
    doc_paths:       list[str],
    output_path:     Path,
    verbose:         bool = False,
    dpi:             int  = 150,
    ocr_dpi:         int  = 200,
    skip_t1:         bool = False,
    skip_paso1:      bool = False,
    skip_paso3:      bool = False,
    skip_paso4_csv:  bool = False,
    skip_hs_ocr:     bool = False,
    skip_validation: bool = False,
) -> int:

    print('\n=== Camión Export Processor (DAE + T1 + validación) ===\n')
    print(f'📂 Input:  {xlsx_path}')
    print(f'💾 Output: {output_path}')

    # ── 1. T1 PDFs ───────────────────────────────────────────────────────────
    t1_info_list: list[dict] = []
    t1_map: dict[str, float] = {}
    if t1_paths:
        print('\n📄 Extracting T1 data...')
        for pdf_path in t1_paths:
            info = t1_processor.extract_t1_info(pdf_path)
            t1_info_list.append(info)
            if info['mrn']:
                t1_map[info['mrn']] = info['gross_kg']
                print(f"  ✓ {info['source_file']}  →  MRN: {info['mrn']}"
                      f" | Gross: {info['gross_kg']:.0f} kg"
                      f" | Pkgs: {info['packages']}"
                      f" | Deadline: {info['deadline']}")
            else:
                print(f"  ⚠ Could not extract MRN from: {pdf_path}")
    else:
        print('\nℹ️  Sin T1 PDFs.')

    # ── 2. DAE_CIERRE_CAMION ─────────────────────────────────────────────────
    src_rows = src_summary = None
    paso1_rows = None
    if not skip_paso1 or not skip_paso3 or not skip_paso4_csv:
        # Lee la fuente una sola vez (la usan DAE_CIERRE_CAMION y DAE_LISTADO)
        src_rows, src_summary = read_source_paso1(xlsx_path)
        print(f'\n📊 Fuente leída: {len(src_rows)} filas')

    if not skip_paso1:
        print('\n── DAE_CIERRE_CAMION (cierre del camión) ─────────────────')
        paso1_rows = process_paso1(src_rows)
        print(f'  ✓ {len(paso1_rows)} filas, Peso Bruto redondeado half-up')

    # ── 3. T1_CIERRE_CAMION ──────────────────────────────────────────────────
    result_t1 = summary_t1 = None
    if not skip_t1:
        print('\n── T1_CIERRE_CAMION (peso bruto del PDF T1) ──────────────')
        rows_t1, summary_t1 = t1_processor.read_sheet1(xlsx_path)
        print(f'  ✓ {len(rows_t1)} filas leídas')
        result_t1 = t1_processor.process_packing_list(rows_t1, t1_map)

    # ── 4. Validación DOC vs XLSX ────────────────────────────────────────────
    report = None
    if doc_paths and not skip_validation:
        if len(doc_paths) == 1:
            print(f'\n🔎 Validating XLSX vs DOC: {doc_paths[0]}')
        else:
            print(f'\n🔎 Validating XLSX vs {len(doc_paths)} DOCs:')
            for d in doc_paths:
                print(f'    - {d}')
        try:
            import camion_pdf_validator as pdf_validator
            report = pdf_validator.validate_xlsx_against_docs(
                xlsx_path=Path(xlsx_path),
                doc_paths=[Path(p) for p in doc_paths],
                dpi=dpi,
                t1_paths=[Path(p) for p in t1_paths],
            )
        except ImportError:
            print('  ⚠ camion_pdf_validator.py no encontrado – validación DOC omitida.')
        except AttributeError:
            # Fallback a la antigua función si el módulo cargado es una versión previa
            print('  ℹ Usando validador antiguo (solo primer DOC). '
                  'Actualiza camion_pdf_validator.py para validar todos los DOCs.')
            report = pdf_validator.validate_xlsx_against_doc(
                xlsx_path=Path(xlsx_path),
                doc_path=Path(doc_paths[0]),
                dpi=dpi,
                t1_paths=[Path(p) for p in t1_paths],
            )

    # ── 5. OCR de HS codes desde DOC.pdf(s) ──────────────────────────────────
    hs_map: dict = {}
    if doc_paths and not skip_hs_ocr and (not skip_paso3 or not skip_paso4_csv):
        print('\n🔎 OCR HS codes desde DOC.pdf...')
        try:
            from hs_extractor import extract_hs_map
            hs_map = extract_hs_map(doc_paths, dpi=ocr_dpi, verbose=verbose)

            mrn_part     = hs_map.get('mrn', {})     if isinstance(hs_map, dict) else {}
            shipper_part = hs_map.get('shipper', {}) if isinstance(hs_map, dict) else {}

            if mrn_part:
                print(f'  ✓ {len(mrn_part)} MRN→HS extraídos directamente:')
                for mrn, hs in sorted(mrn_part.items()):
                    print(f'    {mrn}  →  {hs}')
            if shipper_part:
                print(f'  ✓ {len(shipper_part)} HS extraídos por nombre de shipper'
                      f' (bloques con MRN no legible):')
                for ship, hs in sorted(shipper_part.items()):
                    print(f'    shipper={ship!r}  →  {hs}')
            if not mrn_part and not shipper_part:
                print('  ⚠ No se extrajo ningún HS code (columna quedará vacía)')
        except ImportError:
            print('  ⚠ hs_extractor.py no encontrado – HS codes vacíos.')
        except Exception as e:
            print(f'  ⚠ Error en OCR: {e} – HS codes vacíos.')

    # ── 6. DAE_LISTADO + DAE_VISUAL.csv ──────────────────────────────────────
    paso3_rows = None
    if not skip_paso3 or not skip_paso4_csv:
        paso3_rows = build_paso3_rows(src_rows, hs_map=hs_map)
        print(f'\n── DAE_LISTADO (DAE-only con HS) ─────────────────────────')
        print(f'  ✓ {len(paso3_rows)} filas DAE filtradas')
        if verbose:
            for r in paso3_rows:
                print(f'    {r["mrn"]:<22} HS={r["hs"] or "-":<6} '
                      f'bultos={r["bultos"]:<4} PB={r["peso_bruto"]}')

    # ── 7. Workbook final ────────────────────────────────────────────────────
    print('\n💾 Escribiendo workbook final...')
    wb = openpyxl.load_workbook(xlsx_path)   # conserva la hoja original del cliente

    # Renombra la hoja original del cliente a PL_ORIGINAL para que sea
    # autodescriptiva (es el packing-list que manda Sostmeier). Solo si
    # el workbook tiene una sola hoja inicial — si ya viene con varias,
    # no asumimos cuál es la del PL.
    if len(wb.sheetnames) == 1:
        original_sheet = wb.sheetnames[0]
        if original_sheet != 'PL_ORIGINAL':
            wb[original_sheet].title = 'PL_ORIGINAL'

    if report is not None:
        import camion_pdf_validator as pdf_validator
        pdf_validator.add_validated_sheet(wb, report, validated_sheet_name='VALIDACION_PL_DOC')
        print('  ✓ Hoja: VALIDACION_PL_DOC')

    if paso1_rows is not None:
        add_paso1_sheet(wb, paso1_rows, src_summary, sheet_name='DAE_CIERRE_CAMION')
        print('  ✓ Hoja: DAE_CIERRE_CAMION')

    if result_t1 is not None:
        t1_processor.add_processed_sheet(wb, result_t1, summary_t1, processed_sheet_name='T1_CIERRE_CAMION')
        print('  ✓ Hoja: T1_CIERRE_CAMION')

    if t1_info_list:
        t1_processor.add_t1_summary_sheet(wb, t1_info_list, sheet_name='T1_RESUMEN')
        print('  ✓ Hoja: T1_RESUMEN')

    if not skip_paso3 and paso3_rows is not None:
        add_paso3_sheet(wb, paso3_rows, sheet_name='DAE_LISTADO')
        print('  ✓ Hoja: DAE_LISTADO')

    wb.save(output_path)

    # ── 8. CSV DAE_VISUAL ────────────────────────────────────────────────────
    if not skip_paso4_csv and paso3_rows is not None:
        csv_path = derive_csv_path(output_path)
        write_paso4_csv(csv_path, paso3_rows)
        print(f'  ✓ CSV:  {csv_path}')

    print(f'\n✅ Done!')
    print(f'  ✓ XLSX: {output_path}')
    if not skip_paso4_csv and paso3_rows is not None:
        print(f'  ✓ CSV:  {derive_csv_path(output_path)}')
    print()
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(
        description='Camión Export Processor — DAE + T1 + validación',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Inputs:
  --xlsx PL.xlsx          Packing-list (siempre requerido)
  --doc DOC.pdf [...]     PDFs aduaneros (1 o varios). Necesario para --validate y --hs.
  --t1 T1.pdf [...]       PDFs T1. Si los pasas, activa el procesamiento T1.

Procesos (qué se ejecuta):
  --dae / --no-dae        DAE pipeline (DAE_CIERRE_CAMION + DAE_LISTADO + CSV
                          DAE_VISUAL). Default: ON.
  --hs / --no-hs          OCR de HS codes para DAE_LISTADO/CSV. Default: ON
                          cuando hay --doc.
  --validate              Generar hoja VALIDACION_PL_DOC. Default: OFF
                          (requiere flag explícito).

Coherencia: el script aborta si pides algo que no puede hacer
  - --validate sin --doc          (no hay nada contra qué validar)
  - --hs sin --doc                (no hay PDF de donde extraer)
  - --no-dae sin --t1             (no se generaría ningún output)

Ejemplos:

  # Solo PL — rápido, HS y validación se rellenan a mano luego
  run_processors.py --xlsx PL.xlsx

  # PL + DOCs (DAE con HS, sin validación)
  run_processors.py --xlsx PL.xlsx --doc DOC1.pdf DOC2.pdf

  # PL + DOCs + validación
  run_processors.py --xlsx PL.xlsx --doc DOC1.pdf DOC2.pdf --validate

  # PL + DOCs solo para validar, sin OCR de HS (más rápido)
  run_processors.py --xlsx PL.xlsx --doc DOC.pdf --validate --no-hs

  # Pipeline completo: DAE + T1 + validación
  run_processors.py --xlsx PL.xlsx --doc DOC1.pdf DOC2.pdf \\
                    --t1 T1_*.pdf --validate

  # Solo T1 (sin DAE, sin DOCs)
  run_processors.py --xlsx PL.xlsx --t1 T1_*.pdf --no-dae
""",
    )

    # Inputs
    parser.add_argument('--xlsx',         required=True,         help='Packing-list Excel de entrada')
    parser.add_argument('--doc',          nargs='*', default=[], help='DOC PDFs (uno o varios)')
    parser.add_argument('--t1',           nargs='*', default=[], help='T1 PDFs (uno o varios)')
    parser.add_argument('-o', '--output', default=None,          help='Fichero .xlsx o directorio de salida')

    # Procesos: cada uno con su par on/off
    dae_group = parser.add_mutually_exclusive_group()
    dae_group.add_argument('--dae',    dest='dae', action='store_true',  default=True,
                           help='Generar pipeline DAE (CIERRE_CAMION + LISTADO + CSV). [default: ON]')
    dae_group.add_argument('--no-dae', dest='dae', action='store_false',
                           help='No generar pipeline DAE.')

    hs_group = parser.add_mutually_exclusive_group()
    hs_group.add_argument('--hs',    dest='hs', action='store_true',  default=None,
                          help='Extraer HS codes vía OCR del DOC. [default: ON si hay --doc]')
    hs_group.add_argument('--no-hs', dest='hs', action='store_false',
                          help='No extraer HS codes (columna HS queda vacía en DAE_LISTADO/CSV).')

    parser.add_argument('--validate', action='store_true',
                        help='Generar hoja VALIDACION_PL_DOC. Requiere --doc. [default: OFF]')

    parser.add_argument('--verbose',  action='store_true',  help='Log detallado (incluye OCR por página)')

    args = parser.parse_args()

    # ── Resolver defaults dependientes ───────────────────────────────────────
    has_doc = bool(args.doc)
    has_t1  = bool(args.t1)

    # --hs: default ON si hay --doc, OFF si no
    if args.hs is None:
        args.hs = has_doc

    # ── Validaciones de coherencia ───────────────────────────────────────────
    errors = []
    if args.validate and not has_doc:
        errors.append("--validate requiere --doc (no hay PDF contra el que validar)")
    if args.hs and not has_doc:
        errors.append("--hs requiere --doc (no hay PDF de donde extraer HS codes)")
    if not args.dae and not has_t1:
        errors.append("--no-dae sin --t1 no produce ningún output útil")

    if errors:
        print("\n❌ Argumentos incoherentes:", file=sys.stderr)
        for e in errors:
            print(f"   - {e}", file=sys.stderr)
        print("\nEjecuta con --help para ver ejemplos.\n", file=sys.stderr)
        return 2

    # ── Avisos amistosos (no abortan) ────────────────────────────────────────
    if has_doc and not args.validate and not args.hs:
        print("ℹ Has pasado --doc pero no --validate ni --hs activo. "
              "El DOC no se va a usar.", file=sys.stderr)

    # ── Ejecutar ─────────────────────────────────────────────────────────────
    output_path = derive_output_xlsx(args.xlsx, args.output)

    return run(
        xlsx_path       = args.xlsx,
        t1_paths        = args.t1 or [],
        doc_paths       = args.doc or [],
        output_path     = output_path,
        verbose         = args.verbose,
        dpi             = 150,
        ocr_dpi         = 200,
        skip_t1         = not has_t1,
        skip_paso1      = not args.dae,
        skip_paso3      = not args.dae,
        skip_paso4_csv  = not args.dae,
        skip_hs_ocr     = not args.hs,
        skip_validation = not args.validate,
    )


if __name__ == '__main__':
    raise SystemExit(main())