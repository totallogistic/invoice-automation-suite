#!/usr/bin/env python3
"""
processor_paso1.py  –  Generador de la hoja PASO_1 ("cierre del camión")
==============================================================================
Reemplaza a processor_dae.py.

Lee el XLSX fuente (formato Tangier o Kenitra, autodetectando columnas) y
produce la hoja PASO_1 tal y como la describe el cliente:

  - 9 columnas: Shipper | MRN/Invoice | Peso Bruto | Peso Neto |
                PK | BX | PX | CL | RO
  - Sin Tour, Date, Trailer.
  - Peso Bruto redondeado fila a fila con ROUND(x, 0) half-up.
  - Cyan-totales original conservada (peso bruto sin redondear).
  - Alternancia estricta amarillo / blanco en filas de datos.
  - Fila COINCIDE final con SUM real de Peso Bruto y gran total de bultos.

Uso standalone:
    python processor_paso1.py --xlsx PL.xlsx -o PL-PASO1.xlsx

Uso integrado (desde run_processors.py):
    from processor_paso1 import read_source_paso1, process_paso1, add_paso1_sheet
"""

from __future__ import annotations

import math
import argparse
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ─────────────────────────────────────────────────────────────────────────────
# Estilos
# ─────────────────────────────────────────────────────────────────────────────

_GREY_HEADER = 'FFF0F0F0'   # cabeceras + "Name"
_CYAN_TOTALS = 'FFCCFFFF'   # fila de totales originales
_YELLOW_DATA = 'FFFFFFBB'   # alternancia amarilla
# blanco (sin fill) = filas impares de datos


def _fill(hex_rgb: str) -> PatternFill:
    return PatternFill('solid', start_color=hex_rgb, end_color=hex_rgb)


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────

def _round_half_up(x: float) -> int:
    """ROUND(x, 0) equivalente a Excel: 0.5 siempre redondea hacia arriba."""
    return math.floor(x + 0.5)


def _detect_columns(header_row: tuple) -> dict:
    """
    Detecta índices de columnas a partir de la fila de cabeceras.
    Compatible con formato Kenitra (22 cols, sin BX/PX/RO) y
    formato Tangier (25 cols, con BX/PX/RO).
    """
    h = [str(v).strip().upper() if v is not None else '' for v in header_row]

    def find(name: str, start: int = 0) -> Optional[int]:
        for i in range(start, len(h)):
            if h[i] == name.upper():
                return i
        return None

    pb_idx = find('PESO BRUTO')
    pk_idx = find('PK', (pb_idx or 17) + 1)

    return {
        'shipper_name': 4,                                      # posicional, igual en ambos formatos
        'mrn_invoice':  find('MRN / INVOICE') or 13,
        'peso_bruto':   pb_idx if pb_idx is not None else 18,
        'peso_neto':    (pb_idx + 1) if pb_idx is not None else 19,
        'pk':           pk_idx if pk_idx is not None else 20,
        'bx':           find('BX', (pk_idx or 20) + 1),         # None en Kenitra
        'px':           find('PX', (pk_idx or 20) + 1),         # None en Kenitra
        'cl':           find('CL', (pk_idx or 20) + 1),
        'ro':           find('RO', (pk_idx or 20) + 1),         # None en Kenitra
    }


def _get(row: tuple, col: dict, key: str):
    idx = col.get(key)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _find_source_worksheet(wb: openpyxl.Workbook):
    """
    Busca la hoja del packing list por contenido de cabeceras (no por nombre).
    """
    for ws in wb.worksheets:
        try:
            row1 = [str(ws.cell(1, c).value or "").strip().lower() for c in range(1, 16)]
            has_base = (
                len(row1) >= 4
                and row1[0] == "tour"
                and row1[1] == "date"
                and row1[2] == "trailer"
            )
            if has_base:
                return ws
        except Exception:
            continue
    raise ValueError("No se encontró ninguna hoja con formato válido de packing list")


# ─────────────────────────────────────────────────────────────────────────────
# Reader
# ─────────────────────────────────────────────────────────────────────────────

def read_source_paso1(xlsx_path: str) -> tuple[list[dict], dict]:
    """
    Lee la hoja fuente para PASO_1.

    Returns:
        rows    – filas de datos (R5+ del XLSX fuente)
        summary – fila de totales (cyan, normalmente R4)
    """
    wb = openpyxl.load_workbook(xlsx_path)
    ws = _find_source_worksheet(wb)
    raw_rows = list(ws.iter_rows(values_only=False))
    raw      = [tuple(c.value for c in r) for r in raw_rows]

    col = _detect_columns(raw[0])

    # Localiza la fila de totales (primera fila entre 2..5 con peso_bruto)
    summary_row_idx = None
    for i in range(2, min(6, len(raw))):
        if _get(raw[i], col, 'peso_bruto') is not None:
            summary_row_idx = i
            break
    if summary_row_idx is None:
        raise ValueError("No se encontró la fila de totales (cyan) en el XLSX fuente")

    s = raw[summary_row_idx]
    summary = {
        'total_peso_bruto': _get(s, col, 'peso_bruto'),
        'total_peso_neto':  _get(s, col, 'peso_neto'),
        'total_pk':         _get(s, col, 'pk'),
        'total_bx':         _get(s, col, 'bx'),
        'total_px':         _get(s, col, 'px'),
        'total_cl':         _get(s, col, 'cl'),
        'total_ro':         _get(s, col, 'ro'),
    }

    rows = []
    for raw_row in raw[summary_row_idx + 1:]:
        if all(v is None for v in raw_row):
            continue
        rows.append({
            'shipper_name': _get(raw_row, col, 'shipper_name'),
            'mrn_invoice':  _get(raw_row, col, 'mrn_invoice'),
            'peso_bruto':   _get(raw_row, col, 'peso_bruto'),
            'peso_neto':    _get(raw_row, col, 'peso_neto'),
            'pk':           _get(raw_row, col, 'pk'),
            'bx':           _get(raw_row, col, 'bx'),
            'px':           _get(raw_row, col, 'px'),
            'cl':           _get(raw_row, col, 'cl'),
            'ro':           _get(raw_row, col, 'ro'),
        })

    return rows, summary


# ─────────────────────────────────────────────────────────────────────────────
# Transform
# ─────────────────────────────────────────────────────────────────────────────

def process_paso1(rows: list[dict]) -> list[dict]:
    """Aplica ROUND(x, 0) half-up al Peso Bruto de todas las filas."""
    result = []
    for row in rows:
        pb = row['peso_bruto']
        result.append({
            **row,
            'peso_bruto': _round_half_up(pb) if pb is not None else None,
        })
    return result


# ─────────────────────────────────────────────────────────────────────────────
# Sheet writer
# ─────────────────────────────────────────────────────────────────────────────

def add_paso1_sheet(
    wb:          openpyxl.Workbook,
    result_rows: list[dict],
    summary:     dict,
    sheet_name:  str = 'PASO_1',
) -> None:
    """
    Añade la hoja PASO_1 al workbook.

    Layout exacto (replica el PASO_1.xlsx manual del cliente):
        R1: cabeceras (gris)
        R2: "Name" en C1 (gris)
        R3: vacía (sin fill)
        R4: totales originales (cyan)
        R5+: datos – alternancia estricta blanco/amarillo (R5 blanco, R6 amarillo, ...)
        Última-1: vacía (sin fill)
        Última: COINCIDE – C3=SUM(PB), C9=gran total bultos
    """
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    arial10 = Font(name='Arial', size=10)

    headers = ['Shipper', 'MRN / Invoice',
               'Peso Bruto', 'Peso Neto',
               'PK', 'BX', 'PX', 'CL', 'RO']

    # R1 - cabeceras
    ws.append(headers)
    for cell in ws[1]:
        cell.font      = arial10
        cell.fill      = _fill(_GREY_HEADER)
        cell.alignment = Alignment(horizontal='left')

    # R2 - "Name" en C1
    ws.append(['Name'] + [None] * 8)
    for cell in ws[2]:
        cell.font = arial10
        cell.fill = _fill(_GREY_HEADER)

    # R3 - vacía sin fill
    ws.append([None] * 9)

    # R4 - totales originales (cyan)
    ws.append([
        None, None,
        summary['total_peso_bruto'], summary['total_peso_neto'],
        summary['total_pk'] or 0,
        summary['total_bx'] or 0,
        summary['total_px'] or 0,
        summary['total_cl'] or 0,
        summary['total_ro'] or 0,
    ])
    for cell in ws[4]:
        cell.font = arial10
        cell.fill = _fill(_CYAN_TOTALS)

    # R5+ filas de datos con alternancia estricta
    data_start_row = 5
    for i, row in enumerate(result_rows):
        ws.append([
            row['shipper_name'],
            row['mrn_invoice'],
            row['peso_bruto'],
            row['peso_neto'],
            row['pk'],
            row['bx'],
            row['px'],
            row['cl'],
            row['ro'],
        ])
        excel_row_idx = ws.max_row     # ← capturar TRAS el append
        # Alternancia: par→blanco (sin fill), impar→amarillo
        is_yellow = (i % 2) == 1
        for cell in ws[excel_row_idx]:
            cell.font = arial10
            if is_yellow:
                cell.fill = _fill(_YELLOW_DATA)

    data_end_row = ws.max_row

    # Separador vacío
    ws.append([None] * 9)

    # Fila COINCIDE
    coincide_row = ws.max_row + 1
    ws.append([
        None, None,
        f'=SUM(C{data_start_row}:C{data_end_row})',  # SUM de Peso Bruto redondeado
        None, None, None, None, None,
        f'=SUM(E4:I4)',                              # gran total bultos = PK+BX+PX+CL+RO de la cyan
    ])
    for cell in ws[coincide_row]:
        cell.font = arial10

    # Anchos de columna
    widths = {'A': 35, 'B': 32, 'C': 12, 'D': 12,
              'E': 6, 'F': 6, 'G': 6, 'H': 6, 'I': 6}
    for col_letter, w in widths.items():
        ws.column_dimensions[col_letter].width = w

    ws.freeze_panes = 'A5'


# ─────────────────────────────────────────────────────────────────────────────
# CLI standalone
# ─────────────────────────────────────────────────────────────────────────────

def main() -> int:
    parser = argparse.ArgumentParser(description='PASO_1 generator (truck closure)')
    parser.add_argument('--xlsx',         required=True, help='Input packing-list Excel')
    parser.add_argument('-o', '--output', default=None,  help='Output Excel path')
    parser.add_argument('--verbose',      action='store_true')
    args = parser.parse_args()

    output_path = Path(args.output) if args.output else \
        Path(args.xlsx).with_name(Path(args.xlsx).stem + '-PASO1.xlsx')

    print('\n=== PASO_1 generator ===\n')
    print(f'📊 Reading: {args.xlsx}')

    rows, summary = read_source_paso1(args.xlsx)
    print(f'  ✓ {len(rows)} filas de datos')
    print(f'  ✓ Cyan totals: PB={summary["total_peso_bruto"]}  PN={summary["total_peso_neto"]}'
          f'  PK={summary["total_pk"]}  BX={summary["total_bx"]}'
          f'  PX={summary["total_px"]}  CL={summary["total_cl"]}'
          f'  RO={summary["total_ro"]}')

    result = process_paso1(rows)

    if args.verbose:
        print(f'\n  {"Shipper":<32} {"MRN":<32} {"PB_src":>10} {"PB_out":>8}')
        print('  ' + '-' * 88)
        for src, out in zip(rows, result):
            print(f'  {str(src["shipper_name"] or "")[:31]:<32}'
                  f' {str(src["mrn_invoice"] or "")[:31]:<32}'
                  f' {str(src["peso_bruto"]):>10} {str(out["peso_bruto"]):>8}')

    print(f'\n💾 Writing: {output_path}')
    wb = openpyxl.load_workbook(args.xlsx)
    add_paso1_sheet(wb, result, summary)
    wb.save(output_path)

    print('\n✅ Done!\n')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
