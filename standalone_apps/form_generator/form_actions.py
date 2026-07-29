"""Dispatcher config-driven de formularios (motor de pipeline).

Cada form declara en config/tools.yaml (sección forms:) un `pipeline`: una lista
de acciones que se ejecutan en orden. Acciones GENÉRICAS (sin código por form):

    append_excel · save_json · render · convert_pdf · email · download · dedup · track · custom

Los renderers (layout bespoke) y handlers (operativa entera) custom se registran
por nombre con @register_renderer / @register_handler. app.py inyecta sus helpers
con configure() para no duplicar lógica.

Un paso puede ser una cadena ("append_excel") o un dict con params
({"append_excel": {"rows_from": "gastos"}}). Con `when: <campo>` un paso solo se
ejecuta si ese campo del payload es truthy.
"""
from __future__ import annotations

import base64
import fcntl
import json
import os
import smtplib
import subprocess
from datetime import datetime
from email import encoders
from email.mime.base import MIMEBase
from email.mime.image import MIMEImage
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from pathlib import Path

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

# ── Registries ────────────────────────────────────────────────────────────────
_RENDERERS: dict = {}
_HANDLERS: dict = {}


def register_renderer(name):
    def deco(fn):
        _RENDERERS[name] = fn
        return fn
    return deco


def register_handler(name):
    def deco(fn):
        _HANDLERS[name] = fn
        return fn
    return deco


# ── Config inyectada desde app.py ─────────────────────────────────────────────
_D: dict = {}


def configure(**deps):
    """Inyecta dependencias de app.py: output_dir, excel_dir, get_form_output_dir."""
    _D.update(deps)


def _out_dir(schema_name: str) -> Path:
    fn = _D.get("get_form_output_dir")
    return fn(schema_name) if fn else _D.get("excel_dir", Path("."))


# ── Utilidades ────────────────────────────────────────────────────────────────
def _normalize(step):
    if isinstance(step, str):
        return step, {}
    if isinstance(step, dict) and len(step) == 1:
        k = next(iter(step))
        return k, (step[k] or {})
    raise HTTPException(500, f"Paso de pipeline inválido: {step!r}")


def _fmt(tpl: str, data: dict) -> str:
    try:
        return tpl.format(**data)
    except Exception:
        return tpl


def _when_ok(params: dict, payload: dict) -> bool:
    cond = params.get("when")
    return True if not cond else bool(payload.get(cond))


def _ts() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


# ── Modo de envío (postfix | custom) ──────────────────────────────────────────
def _smtp_settings(mode=None):
    """(host, port, user, pwd, mail_from) según el MODO DE ENVÍO.
    mode: 'postfix' (relay local, sin auth) | 'custom' (SMTP externo autenticado).
    None → MAIL_SEND_MODE (def 'postfix'). Retrocompat: cae a SMTP_* planos."""
    mode = (mode or os.getenv("MAIL_SEND_MODE", "postfix") or "postfix").lower()
    mail_from = os.getenv("MAIL_FROM", "")
    if mode == "custom":
        host = os.getenv("SMTP_CUSTOM_HOST") or os.getenv("SMTP_HOST", "")
        port = int(os.getenv("SMTP_CUSTOM_PORT") or os.getenv("SMTP_PORT") or "587")
        user = os.getenv("SMTP_CUSTOM_USER") or os.getenv("SMTP_USER", "")
        pwd = os.getenv("SMTP_CUSTOM_PASS") or os.getenv("SMTP_PASS", "")
    else:  # postfix
        host = os.getenv("SMTP_POSTFIX_HOST") or os.getenv("SMTP_HOST", "") or "172.18.0.1"
        port = int(os.getenv("SMTP_POSTFIX_PORT") or os.getenv("SMTP_PORT") or "25")
        user = os.getenv("SMTP_POSTFIX_USER", "")
        pwd = os.getenv("SMTP_POSTFIX_PASS", "")
    return host, port, user, pwd, mail_from


def _resolve_mode(ctx, p):
    """Modo efectivo: param de la acción > cfg del form > MAIL_SEND_MODE global."""
    return (p.get("mode_send_email") or p.get("mode")
            or (ctx.get("cfg") or {}).get("mode_send_email") or None)


# ── Email genérico ────────────────────────────────────────────────────────────
def _send_email(to_email, subject, body_text, attach_path=None, mode=None, mail_from=None) -> bool:
    host, port, user, pwd, default_from = _smtp_settings(mode)
    if not host or not to_email:
        return False
    mail_from = mail_from or default_from or user or "procesos@totallogistic.es"
    msg = MIMEMultipart()
    msg["From"], msg["To"], msg["Subject"] = mail_from, to_email, subject
    msg.attach(MIMEText(body_text, "plain", "utf-8"))
    if attach_path:
        p = Path(attach_path)
        subtype = {"pdf": "pdf", "json": "json"}.get(p.suffix.lstrip("."), "octet-stream")
        part = MIMEBase("application", subtype)
        part.set_payload(p.read_bytes())
        encoders.encode_base64(part)
        part.add_header("Content-Disposition", f'attachment; filename="{p.name}"')
        msg.attach(part)
    try:
        if port == 465:
            with smtplib.SMTP_SSL(host, port) as s:
                if user and pwd:
                    s.login(user, pwd)
                s.send_message(msg)
        else:
            with smtplib.SMTP(host, port) as s:
                if user and pwd:
                    s.starttls()
                    s.login(user, pwd)
                s.send_message(msg)
        return True
    except Exception as e:
        print(f"[form_actions] email error: {e}")
        return False


# ── Email del QR/enlace al camionero (HTML + QR incrustado) ───────────────────
_TIPO_LABEL = {
    "deca": "DeCA · Documento de control de transporte",
    "carta_porte": "Carta de porte",
    "cmr": "CMR · Carta de porte internacional",
}


def _send_qr_email(to_email, subject, docs, payload, mode=None, mail_from=None) -> bool:
    """Envía al camionero un email con el QR + enlace de CADA documento generado
    (DeCA, carta de porte, CMR…). `docs` es la lista de dicts {tipo, url, qr_data_uri}.
    Respeta el MODO DE ENVÍO (postfix | custom); ver _smtp_settings."""
    host, port, user, pwd, default_from = _smtp_settings(mode)
    if not host or not to_email or not docs:
        return False
    mail_from = mail_from or default_from or user or "almacenalgeciras@totallogistic.es"
    ruta = " → ".join(x for x in [payload.get("origen"), payload.get("destino")] if x)
    matricula = payload.get("matricula_tractora", "") or ""
    fecha = payload.get("fecha_transporte", "") or ""

    root = MIMEMultipart("related")
    root["From"], root["To"], root["Subject"] = mail_from, to_email, subject
    alt = MIMEMultipart("alternative")
    root.attach(alt)

    # ── Texto plano (todos los documentos) ──
    txt = [f"Documentos de transporte — matrícula {matricula} · {fecha} · {ruta}".strip(" —·"), ""]
    for d in docs:
        txt.append(f"- {_TIPO_LABEL.get(d.get('tipo'), d.get('tipo') or 'Documento')}: {d.get('url')}")
    txt += ["", "Escanea cada QR o abre su enlace. El DeCA debe poder mostrarse en un control "
            "de carretera.", "", "Total Logistic Services, S.L."]
    alt.attach(MIMEText("\n".join(txt), "plain", "utf-8"))

    # ── HTML: una tarjeta (QR + enlace) por documento ──
    meta = "".join(
        f'<tr><td style="padding:2px 10px 2px 0;color:#6b7280;font-size:13px">{k}</td>'
        f'<td style="padding:2px 0;color:#111827;font-size:13px;font-weight:600">{v}</td></tr>'
        for k, v in [("Matrícula", matricula), ("Fecha", fecha), ("Ruta", ruta)] if v
    )
    images, cards = [], []
    for i, d in enumerate(docs):
        cid = f"qr{i}"
        label = _TIPO_LABEL.get(d.get("tipo"), d.get("tipo") or "Documento")
        url = d.get("url") or ""
        cards.append(
            f'<div style="border:1px solid #e5e7eb;border-radius:10px;padding:16px;margin:0 0 14px">'
            f'<div style="font-weight:700;color:#1a4d7e;font-size:15px;margin-bottom:10px">{label}</div>'
            f'<div style="text-align:center;margin-bottom:10px"><img src="cid:{cid}" width="150" height="150" '
            f'style="border:1px solid #e5e7eb;border-radius:8px;padding:6px;background:#fff"></div>'
            f'<div style="text-align:center"><a href="{url}" style="display:inline-block;background:#1a4d7e;'
            f'color:#fff;text-decoration:none;padding:9px 20px;border-radius:8px;font-weight:600;font-size:14px">'
            f'Abrir {label}</a></div>'
            f'<div style="font-size:11px;color:#9ca3af;word-break:break-all;margin-top:8px;text-align:center">{url}</div>'
            f'</div>'
        )
        qr = d.get("qr_data_uri")
        if qr and "," in qr:
            try:
                img = MIMEImage(base64.b64decode(qr.split(",", 1)[1]))
                img.add_header("Content-ID", f"<{cid}>")
                img.add_header("Content-Disposition", f"inline; filename={cid}.png")
                images.append(img)
            except Exception:
                pass
    html = f"""<div style="font-family:-apple-system,Segoe UI,Roboto,sans-serif;max-width:520px;margin:0 auto">
  <div style="background:linear-gradient(135deg,#2563a8,#1a4d7e);color:#fff;padding:20px 24px;border-radius:12px 12px 0 0">
    <div style="font-size:12px;letter-spacing:1px;opacity:.85;text-transform:uppercase">Total Logistic Services, S.L.</div>
    <div style="font-size:19px;font-weight:700;margin-top:2px">Documentos de transporte ({len(docs)})</div>
  </div>
  <div style="padding:22px 24px;border:1px solid #e5e7eb;border-top:none;border-radius:0 0 12px 12px">
    <table style="border-collapse:collapse;margin-bottom:16px">{meta}</table>
    {''.join(cards)}
    <p style="font-size:12px;color:#9ca3af;margin:6px 0 0;line-height:1.5">Escanea cada QR o abre su enlace. Debes poder mostrar el <b>DeCA</b> en un control de carretera.</p>
  </div>
</div>"""
    alt.attach(MIMEText(html, "html", "utf-8"))
    for img in images:
        root.attach(img)
    try:
        if port == 465:
            with smtplib.SMTP_SSL(host, port) as s:
                if user and pwd:
                    s.login(user, pwd)
                s.send_message(root)
        else:
            with smtplib.SMTP(host, port) as s:
                if user and pwd:
                    s.starttls()
                    s.login(user, pwd)
                s.send_message(root)
        return True
    except Exception as e:
        print(f"[form_actions] email_qr error: {e}")
        return False


# ── Store JSON con flock (dedup / track) ──────────────────────────────────────
def _store_path(ctx, params) -> Path:
    return _out_dir(ctx["schema_name"]) / params.get("store", "_processed.json")


def _load_store(path: Path) -> dict:
    if not path.exists():
        return {}
    try:
        with open(path) as f:
            fcntl.flock(f.fileno(), fcntl.LOCK_SH)
            try:
                return json.load(f) or {}
            finally:
                fcntl.flock(f.fileno(), fcntl.LOCK_UN)
    except Exception:
        return {}


def _append_store(path: Path, key: str, entry: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    mode = "r+" if path.exists() else "w+"
    with open(path, mode) as f:
        fcntl.flock(f.fileno(), fcntl.LOCK_EX)
        try:
            data = {}
            if mode == "r+":
                f.seek(0)
                try:
                    data = json.load(f) or {}
                except json.JSONDecodeError:
                    data = {}
            data[key] = entry
            f.seek(0)
            f.truncate()
            json.dump(data, f, indent=2, ensure_ascii=False)
        finally:
            fcntl.flock(f.fileno(), fcntl.LOCK_UN)


# ── Acciones ──────────────────────────────────────────────────────────────────
def _a_append_excel(ctx, p):
    payload, schema = ctx["payload"], ctx["schema_name"]
    out = _out_dir(schema) / f"{p.get('file', schema)}.xlsx"
    base = {k: v for k, v in payload.items() if not isinstance(v, (list, dict))}
    rows_from = p.get("rows_from")
    if rows_from:
        items = payload.get(rows_from, []) or []
        rows = [{**base, **(it if isinstance(it, dict) else {rows_from: it})} for it in items]
    else:
        rows = [base]
    if not rows:
        return
    if out.exists():
        wb = load_workbook(out)
        ws = wb.active
        headers = [c.value for c in ws[1]]
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = schema[:31]
        headers = list(rows[0].keys())
        ws.append(headers)
        fill = PatternFill(start_color="1a4d7e", end_color="1a4d7e", fill_type="solid")
        for c in ws[1]:
            c.fill = fill
            c.font = Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    wb.save(out)
    ctx["response"]["excel_file"] = out.name
    ctx["response"]["rows_added"] = len(rows)


def _a_save_json(ctx, p):
    path = _D["output_dir"] / f"{ctx['schema_name']}_{_ts()}.json"
    path.write_text(json.dumps(ctx["payload"], indent=2, ensure_ascii=False), encoding="utf-8")
    ctx["last_file"] = path
    ctx["response"]["json_file"] = path.name


def _a_render(ctx, p):
    r = _RENDERERS.get(p.get("renderer"))
    if not r:
        raise HTTPException(500, f"Renderer no registrado: {p.get('renderer')}")
    path = Path(r(ctx["payload"], p.get("params", {}), ctx))
    ctx["last_file"] = path
    ctx["response"]["file"] = path.name


def _a_convert_pdf(ctx, p):
    src = ctx.get("last_file")
    if not src:
        raise HTTPException(500, "convert_pdf sin fichero previo (falta un render/save antes)")
    outdir = Path(src).parent
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(src)],
        capture_output=True, text=True, timeout=90, check=True,
    )
    pdf = Path(src).with_suffix(".pdf")
    if not pdf.exists():
        raise HTTPException(500, "PDF no generado por LibreOffice")
    ctx["last_file"] = pdf
    ctx["response"]["pdf_file"] = pdf.name


def _a_email(ctx, p):
    schema = ctx["schema_name"]
    key = schema.upper().replace("-", "_")
    to = (p.get("to")
          or (os.getenv(p["to_env"], "").strip() if p.get("to_env") else "")
          or os.getenv(f"MAIL_TO_{key}", "").strip()
          or os.getenv("MAIL_TO", ""))
    attach = p.get("attach", "json")
    if attach == "last":
        attach_path = ctx.get("last_file")
    elif attach == "json":
        attach_path = _D["output_dir"] / f"{schema}_{_ts()}.json"
        attach_path.write_text(json.dumps(ctx["payload"], indent=2, ensure_ascii=False), encoding="utf-8")
    else:
        attach_path = None
    subject = _fmt(p.get("subject", f"Formulario: {schema}"), ctx["payload"])
    body = p.get("body", "Mensaje generado automáticamente por el sistema de formularios de Totallogistic.")
    sent = _send_email(to, subject, body, attach_path, mode=_resolve_mode(ctx, p), mail_from=p.get("mail_from"))
    ctx["response"]["email_sent"] = sent
    ctx["response"]["email_to"] = to if sent else None


def _a_email_qr(ctx, p):
    """Envía al camionero el QR + enlace de TODOS los documentos generados por el
    handler (DeCA, carta de porte, CMR), una tarjeta por documento. Destino
    dinámico por envío (no hay MAIL_TO_ fijo); la descarga local sigue disponible.
    Params: to_field (def email_camionero) · to · subject · mail_from · mode_send_email · when."""
    payload, resp = ctx["payload"], ctx["response"]
    to = (p.get("to") or (payload.get(p.get("to_field", "email_camionero")) or "").strip())
    if not to:
        resp["email_qr_sent"] = False
        return
    docs = list(resp.get("documentos") or [])
    if not docs and resp.get("url"):  # fallback: emisión de un solo documento
        docs = [{"tipo": payload.get("tipo_documento", "deca"),
                 "url": resp.get("url"), "qr_data_uri": resp.get("qr_data_uri")}]
    docs = [d for d in docs if d.get("url")]
    if not docs:
        resp["email_qr_sent"] = False
        resp["email_qr_error"] = "sin documentos generados"
        return
    subject = _fmt(p.get("subject", "Documentos de transporte {matricula_tractora}"), payload)
    sent = _send_qr_email(to, subject, docs, payload, mode=_resolve_mode(ctx, p), mail_from=p.get("mail_from"))
    resp["email_qr_sent"] = sent
    resp["email_qr_to"] = to if sent else None
    resp["email_qr_count"] = len(docs) if sent else 0


def _a_download(ctx, p):
    f = ctx.get("last_file")
    if f:
        ctx["response"]["download_url"] = f"/download/{Path(f).name}"


def _a_dedup(ctx, p):
    key = str(ctx["payload"].get(p.get("key", ""), "")).strip()
    if not key:
        return
    store = _load_store(_store_path(ctx, p))
    if key in store:
        raise HTTPException(409, detail={"error": "Ya procesado anteriormente",
                                         "key": key, "previous": store[key]})


def _a_track(ctx, p):
    key = str(ctx["payload"].get(p.get("key", ""), "")).strip()
    if not key:
        return
    entry = {"timestamp": datetime.now().isoformat()}
    entry.update({k: v for k, v in ctx["response"].items() if isinstance(v, (str, bool, int, float))})
    _append_store(_store_path(ctx, p), key, entry)


def _a_custom(ctx, p):
    h = _HANDLERS.get(p.get("handler"))
    if not h:
        raise HTTPException(500, f"Handler no registrado: {p.get('handler')}")
    res = h(ctx)
    if isinstance(res, dict):
        ctx["response"].update(res)


def _a_validate(ctx, p):
    """Valida el payload contra el JSON Schema del form (red de los endpoints viejos)."""
    load_schema = _D.get("load_schema")
    if not load_schema:
        return
    from jsonschema import Draft202012Validator
    schema = load_schema(ctx["schema_name"])
    errs = [f"{'/'.join(str(x) for x in e.path) or 'root'}: {e.message}"
            for e in Draft202012Validator(schema).iter_errors(ctx["payload"])]
    if errs:
        raise HTTPException(400, detail={"error": "Validación fallida", "errors": errs})


_ACTIONS = {
    "validate": _a_validate,
    "append_excel": _a_append_excel,
    "save_json": _a_save_json,
    "render": _a_render,
    "convert_pdf": _a_convert_pdf,
    "email": _a_email,
    "email_qr": _a_email_qr,
    "download": _a_download,
    "dedup": _a_dedup,
    "track": _a_track,
    "custom": _a_custom,
}


# ── Runner ────────────────────────────────────────────────────────────────────
def run_pipeline(schema_name: str, payload: dict, forms_config: dict) -> dict:
    """Ejecuta el pipeline declarado del form. Default: [append_excel]."""
    cfg = forms_config.get(schema_name) or {}
    pipeline = cfg.get("pipeline") or ["append_excel"]
    ctx = {"schema_name": schema_name, "payload": payload, "cfg": cfg,
           "last_file": None, "response": {"success": True}}
    for step in pipeline:
        name, params = _normalize(step)
        if not _when_ok(params, payload):
            continue
        action = _ACTIONS.get(name)
        if not action:
            raise HTTPException(500, f"Acción desconocida en el pipeline: {name}")
        action(ctx, params)
    return ctx["response"]