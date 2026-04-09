"""
Extintores/BIEs/Señalización form endpoints.
Add to app.py:  from app_extintores import router_extintores
                app.include_router(router_extintores)
Templates needed in templates/:  template-extintores-pr.xlsx  /  template-extintores-mlg.xlsx
"""
import os, shutil
from datetime import datetime
from fastapi import HTTPException, APIRouter
from pydantic import BaseModel
from typing import List, Optional
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

# ── Models ────────────────────────────────────────────────────────────────────

class CheckItem(BaseModel):
    id: str
    bien: bool = True
    na: bool = False

class ExtintorRow(BaseModel):
    serie: str
    eficacia: str
    altura_ok: bool = False      # False=NO = correcto (≤1.7m)
    fecha_fabric: Optional[int] = None
    fecha_retimbr: Optional[int] = None
    senal: bool = True
    manguera_bien: bool = True
    precinto: bool = True
    accesibilidad_bien: bool = True
    peso_bien: bool = True
    ubicacion: str = ""
    obs: str = ""

class BIEEntry(BaseModel):
    nombre: str
    fecha_fab: Optional[int] = None
    tipo: Optional[int] = None
    checks: List[bool] = []      # True=Bien per trimestral check row

class BIEGenerales(BaseModel):
    senal: bool = True
    accesible: bool = True
    altura_valvula: bool = True
    menos50m: bool = True
    cantidad: bool = True

class PuestoControl(BaseModel):
    numero: int = 1
    presion_suministro: str = "SI"
    presion_sistema: str = "SI"
    apertura_bien: bool = True
    finales_carrera_bien: bool = True
    interruptores_bien: bool = True
    suministro_bien: bool = True
    red_bies: str = ""

class PulsadorEntry(BaseModel):
    nombre: str
    checks: List[bool] = []

class SenalEntry(BaseModel):
    cantidad: Optional[int] = None
    tipo: str = ""
    ubicacion: str = ""
    obs: str = ""

class ExtintoresPayload(BaseModel):
    sede: str
    dia: int
    mes: str
    anio: int
    responsable: str = "MANUEL RODRIGUEZ SANTANA"
    ext_checks: List[CheckItem] = []
    ext_lista: List[ExtintorRow] = []
    ext_conclusion_ok: bool = True
    ext_anomalias: str = ""
    bie_lista: List[BIEEntry] = []
    bie_generales: Optional[BIEGenerales] = None
    bie_puestos: List[PuestoControl] = []
    bie_conclusion_ok: bool = True
    bie_anomalias: str = ""
    sen_checks: List[CheckItem] = []
    sen_lista: List[SenalEntry] = []
    sen_conclusion_ok: bool = True
    pul_lista: List[PulsadorEntry] = []
    pul_conclusion_ok: bool = True

# ── Helpers ───────────────────────────────────────────────────────────────────

def _w(ws, row, col, value):
    cell = ws.cell(row=row, column=col)
    if not isinstance(cell, MergedCell):
        cell.value = value

def _x(b): return "X" if b else ""
def _sino(b): return "SI" if b else "NO"
def _bm(b):   return "BIEN" if b else "MAL"

def _trimestre(mes):
    m = mes.upper()
    if m in ["ENERO","FEBRERO","MARZO"]:    return 1
    if m in ["ABRIL","MAYO","JUNIO"]:        return 2
    if m in ["JULIO","AGOSTO","SEPTIEMBRE"]: return 3
    return 4

# ── Fill: dates ───────────────────────────────────────────────────────────────

def _dates_pr(d, wb):
    cfg = {
        'Datos Generales':           (49, 7, 15, 20, 27, "PUERTO REAL"),
        'BIEs':                      (119, 8, 16, 21, 28, "PUERTO REAL"),
        'Señalización Luminiscente': (98, 7, 15, 20, 27, "PUERTO REAL"),
        'Extintores_2':              (117, 7, 15, 20, 27, "PUERTO REAL"),
        'PULSADORES':                (120, 7, 15, 20, 27, "PUERTO REAL"),
    }
    for sn, (r, cc, cd, cm, ca, ciudad) in cfg.items():
        if sn not in wb.sheetnames: continue
        ws = wb[sn]
        _w(ws,r,cc,ciudad); _w(ws,r,cd,d.dia); _w(ws,r,cm,d.mes.upper()); _w(ws,r,ca,str(d.anio))

def _dates_mlg(d, wb):
    cfg = {
        'Datos Generales':           (49, 7, 15, 20, 27, "MALAGA"),
        'BIEs - CTM':                (160, 7, 15, 20, 27, "MALAGA"),
        'Extintores':                (183, 7, 15, 20, 27, "MALAGA"),
        'Señalización Luminiscente': (133, 7, 15, 20, 27, "MALAGA"),
    }
    for sn, (r, cc, cd, cm, ca, ciudad) in cfg.items():
        if sn not in wb.sheetnames: continue
        ws = wb[sn]
        _w(ws,r,cc,ciudad); _w(ws,r,cd,d.dia); _w(ws,r,cm,d.mes.upper()); _w(ws,r,ca,str(d.anio))

# ── Fill: extintores sheet ────────────────────────────────────────────────────

_EXT_CHECK_ROWS = {"1.1":65,"1.2":66,"1.3":67,"1.4":68,"1.5":69,"1.6":70,
                   "1.7":71,"1.8":72,"1.9":73,"1.10":74,"1.11":75,"1.12":76,
                   "1.13":77,"1.14":79}

def _fill_ext_checks(d, ws):
    for c in d.ext_checks:
        r = _EXT_CHECK_ROWS.get(c.id)
        if r: _w(ws,r,39,_x(c.bien)); _w(ws,r,41,_x(not c.bien and not c.na))

def _fill_ext_lista_pr(d, ws):
    for i, e in enumerate(d.ext_lista):
        r = 95+i
        _w(ws,r,1,str(e.serie)); _w(ws,r,5,e.eficacia)
        _w(ws,r,8,"NO" if not e.altura_ok else "SI")
        if e.fecha_fabric:  _w(ws,r,11,e.fecha_fabric)
        if e.fecha_retimbr: _w(ws,r,14,e.fecha_retimbr)
        _w(ws,r,17,_sino(e.senal)); _w(ws,r,19,_bm(e.manguera_bien))
        _w(ws,r,22,_sino(e.precinto)); _w(ws,r,25,_bm(e.accesibilidad_bien))
        _w(ws,r,28,_bm(e.peso_bien)); _w(ws,r,33,e.ubicacion)
        if e.obs: _w(ws,r,40,e.obs)
    if d.ext_conclusion_ok: _w(ws,108,6,"X")

def _fill_ext_lista_mlg(d, ws):
    for i, e in enumerate(d.ext_lista):
        r = 87+i
        _w(ws,r,1,str(e.serie)); _w(ws,r,5,e.eficacia)
        _w(ws,r,8,"NO" if not e.altura_ok else "SI")
        if e.fecha_retimbr: _w(ws,r,14,e.fecha_retimbr)
        _w(ws,r,17,_sino(e.senal)); _w(ws,r,19,_bm(e.manguera_bien))
        _w(ws,r,22,_sino(e.precinto)); _w(ws,r,25,_bm(e.accesibilidad_bien))
        _w(ws,r,28,_bm(e.peso_bien)); _w(ws,r,33,e.ubicacion)
        if e.obs: _w(ws,r,40,e.obs)
    if d.ext_conclusion_ok: _w(ws,125,6,"X")
    elif d.ext_anomalias:   _w(ws,127,6,"X"); _w(ws,129,9,d.ext_anomalias)

# ── Fill: BIEs sheet ──────────────────────────────────────────────────────────

_PR_BIE_TROWS  = [87,88,89,90,91,93,94,95]   # 8 T-check rows
_PR_BIE_GROWS  = [102,103,104,105,106]
_MLG_BIE_TROWS = [81,82,83,84,85,86,87,88,89,90,91,94,95,96,97,99,100]  # T/A rows

def _fill_bies_pr(d, ws):
    for bi, b in enumerate(d.bie_lista):
        cb, cm_ = 13+4*bi, 15+4*bi
        _w(ws,82,cb,b.nombre)
        if b.fecha_fab: _w(ws,98,cb,b.fecha_fab)
        if b.tipo:      _w(ws,101,cb,b.tipo)
        for ci,row in enumerate(_PR_BIE_TROWS):
            bien = b.checks[ci] if ci < len(b.checks) else True
            _w(ws,row,cb,_x(bien)); _w(ws,row,cm_,_x(not bien))
    gen = d.bie_generales or BIEGenerales()
    gv = [gen.senal,gen.accesible,gen.altura_valvula,gen.menos50m,gen.cantidad]
    for ri,row in enumerate(_PR_BIE_GROWS):
        _w(ws,row,13,_x(gv[ri])); _w(ws,row,15,_x(not gv[ri]))
    if d.bie_conclusion_ok: _w(ws,115,6,"X")

def _fill_bies_mlg(d, ws):
    for bi, b in enumerate(d.bie_lista):
        cb = 13+4*bi
        _w(ws,76,cb,b.nombre)
        if b.fecha_fab: _w(ws,92,cb,b.fecha_fab)
        if b.tipo:      _w(ws,95,cb,b.tipo)
        for ci,row in enumerate(_MLG_BIE_TROWS):
            bien = b.checks[ci] if ci < len(b.checks) else True
            _w(ws,row,cb,_x(bien))
    for pi,pc in enumerate(d.bie_puestos):
        r = 121+pi
        _w(ws,r,2,pc.numero); _w(ws,r,6,pc.presion_suministro); _w(ws,r,9,pc.presion_sistema)
        _w(ws,r,12,_x(pc.apertura_bien)); _w(ws,r,14,_x(not pc.apertura_bien))
        _w(ws,r,17,_x(pc.finales_carrera_bien)); _w(ws,r,19,_x(not pc.finales_carrera_bien))
        _w(ws,r,22,_x(pc.interruptores_bien)); _w(ws,r,24,_x(not pc.interruptores_bien))
        _w(ws,r,27,_x(pc.suministro_bien)); _w(ws,r,30,_x(not pc.suministro_bien))
        _w(ws,r,33,pc.red_bies)
    if d.bie_conclusion_ok: _w(ws,135,6,"X")

# ── Fill: señalización ────────────────────────────────────────────────────────

def _fill_sen(d, ws, check_rows, start_row, conc_row):
    for ci,c in enumerate(d.sen_checks):
        if ci >= len(check_rows): break
        r = check_rows[ci]
        _w(ws,r,39,_x(c.bien and not c.na))
        _w(ws,r,41,_x(not c.bien and not c.na))
        _w(ws,r,43,_x(c.na))
    for si,s in enumerate(d.sen_lista):
        r = start_row+si
        if s.cantidad is not None: _w(ws,r,1,s.cantidad)
        _w(ws,r,5,s.tipo); _w(ws,r,33,s.ubicacion)
        if s.obs: _w(ws,r,40,s.obs)
    if d.sen_conclusion_ok: _w(ws,conc_row,6,"X")

# ── Fill: pulsadores ──────────────────────────────────────────────────────────

_PR_PUL_ROWS = [89,90,91,92,93,94,95,96]

def _fill_pul_pr(d, ws):
    for pi,p in enumerate(d.pul_lista):
        cb,cm_ = 13+4*pi, 15+4*pi
        _w(ws,84,cb,p.nombre)
        for ci,row in enumerate(_PR_PUL_ROWS):
            bien = p.checks[ci] if ci < len(p.checks) else True
            _w(ws,row,cb,_x(bien)); _w(ws,row,cm_,_x(not bien))
    if d.pul_conclusion_ok: _w(ws,111,6,"X")

# ── Router ────────────────────────────────────────────────────────────────────

router_extintores = APIRouter()

TEMPLATE_DIR = os.path.join(os.path.dirname(__file__), "templates")
OUTPUT_DIR   = os.path.join(os.path.dirname(__file__), "output")

@router_extintores.post("/api/generar-acta-extintores/{sede}")
async def generar_acta(sede: str, data: ExtintoresPayload):
    sede = sede.lower()
    if sede not in ("pr","mlg"):
        raise HTTPException(400, "sede debe ser 'pr' o 'mlg'")

    tpl = os.path.join(TEMPLATE_DIR, f"template-extintores-{sede}.xlsx")
    if not os.path.exists(tpl):
        raise HTTPException(500, f"Template no encontrado: template-extintores-{sede}.xlsx")

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    ts  = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_name = f"Acta_Extintores_{'PR' if sede=='pr' else 'MLG'}_{data.anio}_T{_trimestre(data.mes)}_{ts}.xlsx"
    out_path = os.path.join(OUTPUT_DIR, out_name)
    shutil.copy2(tpl, out_path)

    wb = load_workbook(out_path)

    if sede == "pr":
        _dates_pr(data, wb)
        _fill_ext_checks(data, wb['Extintores_2'])
        _fill_ext_lista_pr(data, wb['Extintores_2'])
        _fill_bies_pr(data, wb['BIEs'])
        _fill_sen(data, wb['Señalización Luminiscente'],
                  [57,58,59,60,61,62,63,64,65,66], 73, 89)
        if 'PULSADORES' in wb.sheetnames:
            _fill_pul_pr(data, wb['PULSADORES'])
    else:
        _dates_mlg(data, wb)
        _fill_ext_checks(data, wb['Extintores'])
        _fill_ext_lista_mlg(data, wb['Extintores'])
        _fill_bies_mlg(data, wb['BIEs - CTM'])
        _fill_sen(data, wb['Señalización Luminiscente'],
                  [64,65,66,67,68,69,70,71,72,73], 86, 124)

    wb.save(out_path)
    return {"status":"ok","filename":out_name,"download_url":f"/download/{out_name}"}