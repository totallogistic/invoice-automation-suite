"""
Generador de Situación Financiera.

Lee:
  situacion_config.yaml  → todos los datos excepto Dispuesto
  saldos_actuales.yaml   → Dispuesto (K) por IBAN

Escribe sobre el template limpio (solo fórmulas) y genera
data/situacion_financiera.xlsx listo para usar.
"""

import logging
import yaml
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

logger = logging.getLogger(__name__)

BASE_DIR      = Path(__file__).parent
TEMPLATE_PATH = BASE_DIR / "data" / "SITUACION_FINANCIERA_TEMPLATE.xlsx"
OUTPUT_PATH   = BASE_DIR / "data" / "situacion_financiera.xlsx"
CONFIG_PATH   = BASE_DIR / "situacion_config.yaml"
SALDOS_PATH   = BASE_DIR / "data" / "saldos_actuales.yaml"

# Columnas del template
COL = {
    "banco":       1,   # A
    "cta_aux":     2,   # B
    "producto":    3,   # C
    "or":          4,   # D
    "tipo":        5,   # E
    "euribor_lbl": 6,   # F
    "diferencial": 7,   # G
    "nd":          8,   # H
    "importe":     10,  # J
    "dispuesto":   11,  # K
    "reservado":   18,  # R
    "label_tabla": 19,  # S  (etiqueta en tabla R)
}


def _load_yaml(path: Path) -> dict:
    if not path.exists():
        return {}
    data = yaml.safe_load(path.read_text())
    return data if isinstance(data, dict) else {}


def _iban(value) -> str:
    return str(value).replace(" ", "").upper() if value else ""


async def generate_situacion() -> Path:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"Template no encontrado: {TEMPLATE_PATH}\n"
            "Copia SITUACION_FINANCIERA_TEMPLATE.xlsx a la carpeta data/."
        )

    config = _load_yaml(CONFIG_PATH)
    saldos = _load_yaml(SALDOS_PATH)

    if not config:
        raise ValueError("situacion_config.yaml vacío o no encontrado.")

    # Índice IBAN → dispuesto desde saldos_actuales.yaml
    iban_index: dict[str, float] = {}
    for s in saldos.get("saldos", []):
        iban_key = _iban(s.get("iban", ""))
        if iban_key:
            iban_index[iban_key] = float(s.get("dispuesto", 0))

    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active

    # ── 1. Fecha ──────────────────────────────────────────────────────
    ws["F1"] = f"SITUACIÓN FINANCIERA {datetime.now().strftime('%d/%m/%y')}"

    # ── 2. Euribor ────────────────────────────────────────────────────
    eur = config.get("euribor", {})
    if eur.get("Euribor 3"):  ws["B19"] = eur["Euribor 3"]
    if eur.get("Euribor 6"):  ws["B20"] = eur["Euribor 6"]
    if eur.get("Euribor 12"): ws["B21"] = eur["Euribor 12"]

    # ── 3. R8 (disponible importación) ───────────────────────────────
    if config.get("disponible_importacion"):
        ws["R8"] = config["disponible_importacion"]

    # ── 4. Anticipos (filas 23-24) ────────────────────────────────────
    for i, ant in enumerate(config.get("anticipos", []), start=23):
        ws.cell(row=i, column=1).value = ant.get("banco", "")
        ws.cell(row=i, column=2).value = ant.get("numero")

    # ── 5. Cuentas ────────────────────────────────────────────────────
    cuentas = config.get("cuentas") or []
    actualizados = 0
    sin_saldo = []

    for cfg in cuentas:
        fila = cfg.get("fila")
        if not fila:
            continue

        # Escribe todos los campos del config en su columna
        ws.cell(row=fila, column=COL["banco"]).value    = cfg.get("banco")
        ws.cell(row=fila, column=COL["cta_aux"]).value  = cfg.get("cta_auxiliar")
        ws.cell(row=fila, column=COL["or"]).value       = cfg.get("or")
        ws.cell(row=fila, column=COL["tipo"]).value     = cfg.get("tipo")
        ws.cell(row=fila, column=COL["producto"]).value = cfg.get("producto")

        if cfg.get("euribor_tipo"):
            ws.cell(row=fila, column=COL["euribor_lbl"]).value = cfg["euribor_tipo"]

        if cfg.get("diferencial") is not None:
            ws.cell(row=fila, column=COL["diferencial"]).value = cfg["diferencial"]

        if cfg.get("nd"):
            ws.cell(row=fila, column=COL["nd"]).value = cfg["nd"]

        if cfg.get("importe") is not None:
            cell = ws.cell(row=fila, column=COL["importe"])
            cell.value        = cfg["importe"]
            cell.number_format = '#,##0'

        # Reservado (R) — solo pólizas
        if cfg.get("reservado") is not None:
            cell = ws.cell(row=fila, column=COL["reservado"])
            cell.value        = cfg["reservado"]
            cell.number_format = '#,##0'
            # Etiqueta en S (tabla R)
            ws.cell(row=fila, column=COL["label_tabla"]).value = cfg.get("banco")

        # Dispuesto (K) — desde saldos_actuales.yaml por IBAN
        iban = _iban(cfg.get("iban"))
        if iban:
            dispuesto = iban_index.get(iban)
            if dispuesto is not None:
                cell = ws.cell(row=fila, column=COL["dispuesto"])
                cell.value        = round(dispuesto, 2)
                cell.number_format = '#,##0.00'
                actualizados += 1
                logger.info("F%d %s → K=%.2f €", fila, cfg.get("banco", ""), dispuesto)
            else:
                sin_saldo.append(f"  F{fila} {cfg.get('banco')} — IBAN {iban[:16]}... sin datos en saldos_actuales.yaml")

    if sin_saldo:
        logger.info("Cuentas sin Dispuesto:\n%s", "\n".join(sin_saldo))
    logger.info("Dispuesto actualizado: %d cuentas", actualizados)

    # ── 6. Guarda ─────────────────────────────────────────────────────
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    logger.info("Guardado: %s", OUTPUT_PATH)
    return OUTPUT_PATH