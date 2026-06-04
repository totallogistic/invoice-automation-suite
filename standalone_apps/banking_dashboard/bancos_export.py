"""
Generador de Bancos Excel — replica el modelo manual bancos_.xlsx.

Lee:
  config/bancos_config.yaml   → mapeo de cuentas (titular, tipo, parámetros)
  data/saldos_actuales.yaml   → saldos reales del sync bancario

Genera:
  data/bancos_YYYYMMDD.xlsx   → Excel con:
    - Hoja "Saldos actuales": lista plana + secciones por empresa
    - Hoja "Histórico": evolución diaria por banco
"""

import logging
import yaml
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

BASE_DIR     = Path(__file__).parent
CONFIG_PATH  = BASE_DIR / "config" / "bancos_config.yaml"
SALDOS_PATH  = BASE_DIR / "data" / "saldos_actuales.yaml"

# Colores
C_HEADER_BG  = "1A1A2E"
C_HEADER_FG  = "FFFFFF"
C_TOTAL_BG   = "E8F0FE"
C_TITULAR_BG = "2D3748"
C_TITULAR_FG = "FFFFFF"
C_SECTION_BG = "EDF2F7"
C_ALT_ROW    = "F7F9FC"
C_WARN       = "FFF3CD"
C_NEG        = "FFF0F0"
C_BORDER     = "D0D7E3"


def _load_yaml(path: Path) -> dict:
    if not path.exists():
        return {}
    d = yaml.safe_load(path.read_text())
    return d if isinstance(d, dict) else {}


def _iban_key(v) -> str:
    return str(v).replace(" ", "").upper() if v else ""


def _border_thin():
    s = Side(style="thin", color=C_BORDER)
    return Border(left=s, right=s, top=s, bottom=s)


def _header_cell(cell, text, bg=C_HEADER_BG, fg=C_HEADER_FG, bold=True, size=10):
    cell.value = text
    cell.font = Font(name="Arial", bold=bold, color=fg, size=size)
    cell.fill = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _data_cell(cell, value, align="left", bold=False, bg=None, number_fmt=None, size=9):
    cell.value = value
    cell.font = Font(name="Arial", bold=bold, size=size)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = _border_thin()
    if bg:
        cell.fill = PatternFill("solid", start_color=bg)
    if number_fmt:
        cell.number_format = number_fmt


def _currency(v):
    """Formatea para columna de saldo."""
    return '#,##0.00'


async def generate_bancos(history: list[dict] = None, filename: str = None) -> Path:
    config = _load_yaml(CONFIG_PATH)
    saldos = _load_yaml(SALDOS_PATH)

    if not config:
        raise FileNotFoundError(f"bancos_config.yaml no encontrado en {CONFIG_PATH}")

    today = datetime.now()
    if not filename:
        filename = f"bancos_{today.strftime('%Y%m%d')}.xlsx"
    output_path = BASE_DIR / "data" / filename

    # Índice IBAN → saldo del sync
    iban_index: dict[str, dict] = {}
    for s in saldos.get("saldos", []):
        k = _iban_key(s.get("iban", ""))
        if k:
            iban_index[k] = s

    cuentas_cfg = config.get("cuentas") or []
    orden_titulares = config.get("orden_titulares") or []
    euribor = config.get("euribor") or {}

    # Resuelve saldo de cada cuenta
    def get_dispuesto(cfg: dict) -> float:
        iban = _iban_key(cfg.get("iban"))
        if iban and iban in iban_index:
            return float(iban_index[iban].get("dispuesto", 0))
        return float(cfg.get("dispuesto_manual") or cfg.get("dispuesto") or 0)

    def get_moneda(cfg: dict) -> str:
        iban = _iban_key(cfg.get("iban"))
        if iban and iban in iban_index:
            return iban_index[iban].get("moneda", "EUR")
        return "USD" if cfg.get("tipo") == "CTA_CTE_USD" else "EUR"

    wb = Workbook()

    # ──────────────────────────────────────────────────────────
    # HOJA 1: Saldos actuales
    # ──────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Saldos actuales"
    ws.sheet_view.showGridLines = False

    # Anchos de columna
    col_widths = {
        "A": 16, "B": 14, "C": 36, "D": 14, "E": 14,
        "F": 4,  "G": 14, "H": 14, "I": 16,
        "J": 10, "K": 14, "L": 28, "M": 12, "N": 10,
        "O": 10, "P": 8,  "Q": 10, "R": 12, "S": 12,
        "T": 10, "U": 12, "V": 10, "W": 10
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # ── Bloque 1: Lista plana ──
    headers_flat = ["Banco", "IBAN", "Cuenta", "Saldo (EUR)", "Última sync",
                    "", "NO TRAER", "TITULAR", "ENTIDAD"]
    for c, h in enumerate(headers_flat, 1):
        _header_cell(ws.cell(row=1, column=c), h)
    ws.row_dimensions[1].height = 18

    cuentas_con_api = [c for c in cuentas_cfg if _iban_key(c.get("iban")) in iban_index]
    total_valid = 0.0
    flat_row = 2

    for cfg in cuentas_con_api:
        iban = _iban_key(cfg.get("iban"))
        saldo_data = iban_index.get(iban, {})
        dispuesto = float(saldo_data.get("dispuesto", 0))
        moneda = saldo_data.get("moneda", "EUR")
        no_traer = cfg.get("no_traer", False)
        cuenta_nombre = saldo_data.get("cuenta", cfg.get("banco", ""))
        iban_short = f"****{iban[-4:]}" if len(iban) >= 4 else iban
        fecha = saldos.get("fecha", today.strftime("%Y-%m-%d"))

        bg = C_WARN if no_traer else (C_ALT_ROW if flat_row % 2 == 0 else "FFFFFF")

        _data_cell(ws.cell(row=flat_row, column=1), cfg.get("banco", ""), bg=bg)
        _data_cell(ws.cell(row=flat_row, column=2), iban_short, bg=bg)
        _data_cell(ws.cell(row=flat_row, column=3), cuenta_nombre, bg=bg)
        _data_cell(ws.cell(row=flat_row, column=4), dispuesto, align="right",
                   bg=bg, number_fmt="#,##0.00")
        _data_cell(ws.cell(row=flat_row, column=5), fecha, align="center", bg=bg)
        if no_traer:
            _data_cell(ws.cell(row=flat_row, column=7), "NO TRAER",
                       bg=C_WARN, bold=True, align="center", size=8)
        _data_cell(ws.cell(row=flat_row, column=8), cfg.get("titular", ""), bg=bg)
        _data_cell(ws.cell(row=flat_row, column=9), cfg.get("banco", ""), bg=bg)

        if not no_traer and moneda == "EUR":
            total_valid += dispuesto
        flat_row += 1

    # Fila total
    total_row = flat_row
    ws.cell(row=total_row, column=1).value = "TOTAL"
    ws.cell(row=total_row, column=1).font = Font(name="Arial", bold=True, size=10)
    ws.cell(row=total_row, column=1).fill = PatternFill("solid", start_color=C_TOTAL_BG)
    tc = ws.cell(row=total_row, column=4)
    tc.value = round(total_valid, 2)
    tc.font = Font(name="Arial", bold=True, size=10)
    tc.fill = PatternFill("solid", start_color=C_TOTAL_BG)
    tc.number_format = "#,##0.00"
    tc.alignment = Alignment(horizontal="right")

    # Generado
    gen_row = total_row + 2
    ws.cell(row=gen_row, column=1).value = f"Generado: {today.strftime('%d/%m/%Y %H:%M')}"
    ws.cell(row=gen_row, column=1).font = Font(name="Arial", size=9, color="888888")

    # ── Bloque 2: Secciones por TITULAR ──
    section_start = gen_row + 4
    current_row = section_start

    # Agrupa cuentas por titular
    by_titular: dict[str, list] = {t: [] for t in orden_titulares}
    for cfg in cuentas_cfg:
        titular = cfg.get("titular", "")
        if titular in by_titular:
            by_titular[titular].append(cfg)

    def write_poliza_table(ws, row, cuentas_poliza, titulo="TOTAL LOGISTIC SERVICES SL"):
        # Cabecera de sección de pólizas
        _header_cell(ws.cell(row=row, column=8), "IBAN", bg=C_SECTION_BG, fg="2D3748", size=8)
        headers_pol = ["ENTIDAD", "TIPO", "Cta.Aux.", "IBAN", "Referencia",
                       "Euribor", "Diferencial", "ND", "Coste",
                       "Límite", "Dispuesto", "%", "Disponible", "%", "OJO"]
        for c, h in enumerate(headers_pol, 9):
            _header_cell(ws.cell(row=row, column=c), h, bg=C_SECTION_BG, fg="2D3748", size=8)
        row += 1

        total_limite = 0.0
        total_dispuesto = 0.0

        for cfg in cuentas_poliza:
            dispuesto = abs(get_dispuesto(cfg))  # siempre positivo
            iban = _iban_key(cfg.get("iban"))
            iban_short = f"****{iban[-4:]}" if len(iban) >= 4 else "—"
            limite = float(cfg.get("importe") or 0)
            euribor_tipo = cfg.get("euribor_tipo", "")
            euribor_val = euribor.get(euribor_tipo, 0)
            diferencial = float(cfg.get("diferencial") or 0)
            nd_val = float(cfg.get("nd") or 0)
            coste = euribor_val + diferencial
            disponible = limite - dispuesto
            pct_disp = dispuesto / limite if limite else 0
            pct_avail = disponible / limite if limite else 0
            ojo = "OJO NEGATIVO" if get_dispuesto(cfg) < 0 else (
                  " OJO" if pct_avail < 0.4 else "")

            bg = "FFFFFF"
            _data_cell(ws.cell(row=row, column=8), iban_short, bg=bg, size=8)
            _data_cell(ws.cell(row=row, column=9), cfg.get("banco", ""), bg=bg, size=8)
            _data_cell(ws.cell(row=row, column=10), cfg.get("tipo", "PÓLIZA"), bg=bg, size=8)
            _data_cell(ws.cell(row=row, column=11), cfg.get("cta_auxiliar", ""), bg=bg, size=8)
            _data_cell(ws.cell(row=row, column=12), iban, bg=bg, size=8)
            _data_cell(ws.cell(row=row, column=13), euribor_tipo, bg=bg, size=8)    # M: label
            _data_cell(ws.cell(row=row, column=14), euribor_val, align="right",     # N: valor numérico
                        bg=bg, number_fmt="0.000%", size=8)
            _data_cell(ws.cell(row=row, column=15), diferencial, align="right",
                       bg=bg, number_fmt="0.00%", size=8)
            _data_cell(ws.cell(row=row, column=16), nd_val, align="right",
                       bg=bg, number_fmt="0.00%", size=8)
            _data_cell(ws.cell(row=row, column=17), coste, align="right",
                       bg=bg, number_fmt="0.00%", size=8)
            _data_cell(ws.cell(row=row, column=18), limite, align="right",
                       bg=bg, number_fmt="#,##0.00", size=8)
            _data_cell(ws.cell(row=row, column=19), dispuesto, align="right",
                       bg=bg, number_fmt="#,##0.00", size=8)
            _data_cell(ws.cell(row=row, column=20), pct_disp, align="right",
                       bg=bg, number_fmt="0.00%", size=8)
            _data_cell(ws.cell(row=row, column=21), disponible, align="right",
                       bg=bg, number_fmt="#,##0.00", size=8)
            _data_cell(ws.cell(row=row, column=22), pct_avail, align="right",
                       bg=bg, number_fmt="0.00%", size=8)
            if ojo:
                ws.cell(row=row, column=23).value = ojo
                ws.cell(row=row, column=23).font = Font(name="Arial", size=8,
                    color="CC0000" if "NEGATIVO" in ojo else "E65C00")

            total_limite    += limite
            total_dispuesto += dispuesto
            row += 1

        # Totales
        bg_tot = C_TOTAL_BG
        tot_disponible = total_limite - total_dispuesto
        tot_pct_d = total_dispuesto / total_limite if total_limite else 0
        tot_pct_a = tot_disponible / total_limite if total_limite else 0
        _data_cell(ws.cell(row=row, column=18), total_limite, align="right",
                   bg=bg_tot, number_fmt="#,##0.00", bold=True, size=9)
        _data_cell(ws.cell(row=row, column=19), total_dispuesto, align="right",
                   bg=bg_tot, number_fmt="#,##0.00", bold=True, size=9)
        _data_cell(ws.cell(row=row, column=20), tot_pct_d, align="right",
                   bg=bg_tot, number_fmt="0.00%", size=9)
        _data_cell(ws.cell(row=row, column=21), tot_disponible, align="right",
                   bg=bg_tot, number_fmt="#,##0.00", bold=True, size=9)
        _data_cell(ws.cell(row=row, column=22), tot_pct_a, align="right",
                   bg=bg_tot, number_fmt="0.00%", size=9)
        row += 1
        return row

    def write_cta_table(ws, row, cuentas_cta, moneda_label="CTA CTE"):
        headers_cta = ["IBAN", "ENTIDAD", "TIPO", "Cta.Aux.", "IBAN completo",
                       "Dispuesto", "Disponible"]
        _header_cell(ws.cell(row=row, column=8), "", bg=C_SECTION_BG, fg="2D3748")
        for c, h in enumerate(headers_cta, 9):
            _header_cell(ws.cell(row=row, column=c), h, bg=C_SECTION_BG, fg="2D3748", size=8)
        row += 1

        total_disp = 0.0
        for cfg in cuentas_cta:
            dispuesto = get_dispuesto(cfg)
            iban = _iban_key(cfg.get("iban", ""))
            iban_short = f"****{iban[-4:]}" if len(iban) >= 4 else "—"
            bg = "FFFFFF"
            _data_cell(ws.cell(row=row, column=8), iban_short, bg=bg, size=8)
            _data_cell(ws.cell(row=row, column=9), cfg.get("banco", ""), bg=bg, size=8)
            _data_cell(ws.cell(row=row, column=10), moneda_label, bg=bg, size=8)
            _data_cell(ws.cell(row=row, column=11), cfg.get("cta_auxiliar", ""), bg=bg, size=8)
            _data_cell(ws.cell(row=row, column=12), iban, bg=bg, size=8)
            _data_cell(ws.cell(row=row, column=19), dispuesto, align="right",
                       bg=bg, number_fmt="#,##0.00", size=8)
            _data_cell(ws.cell(row=row, column=21), dispuesto, align="right",
                       bg=bg, number_fmt="#,##0.00", size=8)
            total_disp += dispuesto
            row += 1

        # Subtotal
        _data_cell(ws.cell(row=row, column=21), total_disp, align="right",
                   bg=C_TOTAL_BG, number_fmt="#,##0.00", bold=True, size=9)
        row += 2
        return row

    for titular in orden_titulares:
        cuentas_t = by_titular.get(titular, [])
        if not cuentas_t:
            continue

        # Cabecera TITULAR
        tc = ws.cell(row=current_row, column=12)
        tc.value = titular
        tc.font = Font(name="Arial", bold=True, size=11, color=C_TITULAR_FG)
        tc.fill = PatternFill("solid", start_color=C_TITULAR_BG)
        tc.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[current_row].height = 20
        current_row += 2

        # Separa por tipo
        polizas = [c for c in cuentas_t
                   if c.get("tipo") == "POLIZA" and not c.get("no_traer")]
        cta_eur = [c for c in cuentas_t
                   if c.get("tipo") == "CTA_CTE" and not c.get("no_traer")]
        cta_usd = [c for c in cuentas_t
                   if c.get("tipo") == "CTA_CTE_USD" and not c.get("no_traer")]
        lineas  = [c for c in cuentas_t
                   if c.get("tipo") in ("LINEA", "ANTIC") and not c.get("no_traer")]

        if polizas:
            # Cabecera tipo
            ws.cell(row=current_row, column=12).value = "PÓLIZAS"
            ws.cell(row=current_row, column=12).font = Font(
                name="Arial", bold=True, size=9, color="4A5568")
            current_row += 1
            current_row = write_poliza_table(ws, current_row, polizas, titular)
            current_row += 1

        if cta_eur:
            ws.cell(row=current_row, column=12).value = "CTA CTE €"
            ws.cell(row=current_row, column=12).font = Font(
                name="Arial", bold=True, size=9, color="4A5568")
            current_row += 1
            current_row = write_cta_table(ws, current_row, cta_eur, "CTA CTE")

        if cta_usd:
            ws.cell(row=current_row, column=12).value = "CTA CTE $"
            ws.cell(row=current_row, column=12).font = Font(
                name="Arial", bold=True, size=9, color="4A5568")
            current_row += 1
            current_row = write_cta_table(ws, current_row, cta_usd, "CTA CTE $")

        if lineas:
            ws.cell(row=current_row, column=12).value = "LÍNEAS"
            ws.cell(row=current_row, column=12).font = Font(
                name="Arial", bold=True, size=9, color="4A5568")
            current_row += 1
            current_row = write_poliza_table(ws, current_row, lineas, titular)
            current_row += 1

        current_row += 2

    # ──────────────────────────────────────────────────────────
    # HOJA 2: Histórico
    # ──────────────────────────────────────────────────────────
    if history:
        ws2 = wb.create_sheet("Histórico")
        ws2.sheet_view.showGridLines = False

        bank_names = sorted(set(r["bank_name"] for r in history))
        all_days   = sorted(set(r["day"] for r in history))
        day_map    = {}
        for r in history:
            day_map.setdefault(r["day"], {})[r["bank_name"]] = r["total"]

        headers = ["Fecha"] + bank_names + ["Total EUR"]
        for c, h in enumerate(headers, 1):
            _header_cell(ws2.cell(row=1, column=c), h)
        ws2.column_dimensions["A"].width = 14
        for i in range(2, len(headers) + 1):
            ws2.column_dimensions[get_column_letter(i)].width = 16

        for r, day in enumerate(all_days, 2):
            bg = C_ALT_ROW if r % 2 == 0 else "FFFFFF"
            ws2.cell(row=r, column=1).value = day
            ws2.cell(row=r, column=1).alignment = Alignment(horizontal="center")
            row_total = 0.0
            for c, bank in enumerate(bank_names, 2):
                val = day_map.get(day, {}).get(bank)
                cell = ws2.cell(row=r, column=c)
                cell.value = round(val, 2) if val is not None else None
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
                cell.fill = PatternFill("solid", start_color=bg)
                if val:
                    row_total += val
            tot_cell = ws2.cell(row=r, column=len(headers))
            tot_cell.value = round(row_total, 2)
            tot_cell.number_format = "#,##0.00"
            tot_cell.font = Font(name="Arial", bold=True, size=9)
            tot_cell.fill = PatternFill("solid", start_color=C_TOTAL_BG)
            tot_cell.alignment = Alignment(horizontal="right")

    # Guarda
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    logger.info("Bancos Excel guardado: %s", output_path)
    return output_path