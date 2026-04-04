"""
Genera data/saldos.xlsx con dos pestañas:
  - Saldos actuales: último saldo por cuenta
  - Histórico: evolución diaria por banco
Se llama al final de cada sync y desde el endpoint /export/excel.
"""

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

EXCEL_PATH = Path(__file__).parent / "data" / "saldos.xlsx"

# Colores corporativos
COLOR_HEADER_BG  = "1A1A2E"   # azul oscuro del dashboard
COLOR_HEADER_FG  = "FFFFFF"
COLOR_TOTAL_BG   = "E8F0FE"
COLOR_ALT_ROW    = "F7F9FC"
COLOR_POSITIVE   = "1E6823"
COLOR_BORDER     = "D0D7E3"


def _header_style(cell, bg=COLOR_HEADER_BG, fg=COLOR_HEADER_FG, bold=True):
    cell.font = Font(name="Arial", bold=bold, color=fg, size=10)
    cell.fill = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _border():
    side = Side(style="thin", color=COLOR_BORDER)
    return Border(left=side, right=side, top=side, bottom=side)


def _currency_fmt(cell, value):
    cell.value = value
    cell.number_format = '#,##0.00 €'
    cell.font = Font(name="Arial", size=10)
    cell.alignment = Alignment(horizontal="right")


async def generate_excel(latest: list[dict], history: list[dict], daily_totals: list[dict], filename: str = "saldos.xlsx") -> Path:
    wb = Workbook()

    # ------------------------------------------------------------------
    # Pestaña 1: Saldos actuales
    # ------------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Saldos actuales"
    ws1.sheet_view.showGridLines = False

    headers = ["Banco", "IBAN", "Cuenta", "Saldo (EUR)", "Última sync"]
    col_widths = [22, 28, 24, 16, 18]

    for c, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws1.cell(row=1, column=c, value=h)
        _header_style(cell)
        ws1.column_dimensions[get_column_letter(c)].width = w
    ws1.row_dimensions[1].height = 22

    total = 0.0
    for r, row in enumerate(latest, 2):
        bg = COLOR_ALT_ROW if r % 2 == 0 else "FFFFFF"
        fill = PatternFill("solid", start_color=bg)

        iban_display = f"****{row['iban'][-4:]}" if row.get("iban") and len(row["iban"]) >= 4 else row.get("iban", "—")
        sync_date = row["recorded_at"][:10] if row.get("recorded_at") else "—"

        for c, val in enumerate([row["bank_name"], iban_display, row.get("account_name", ""), row["amount"], sync_date], 1):
            cell = ws1.cell(row=r, column=c, value=val)
            cell.fill = fill
            cell.border = _border()
            if c == 4:
                _currency_fmt(cell, val)
                cell.fill = fill
            else:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="center")
        total += row["amount"]

    # Fila total
    total_row = len(latest) + 2
    ws1.cell(row=total_row, column=1, value="TOTAL").font = Font(name="Arial", bold=True, size=10)
    ws1.cell(row=total_row, column=1).fill = PatternFill("solid", start_color=COLOR_TOTAL_BG)
    total_cell = ws1.cell(row=total_row, column=4, value=f"=SUM(D2:D{total_row-1})")
    total_cell.number_format = '#,##0.00 €'
    total_cell.font = Font(name="Arial", bold=True, size=10, color=COLOR_POSITIVE)
    total_cell.fill = PatternFill("solid", start_color=COLOR_TOTAL_BG)
    total_cell.alignment = Alignment(horizontal="right")

    # Metadatos
    ws1.cell(row=total_row + 2, column=1, value=f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}").font = Font(name="Arial", size=9, color="888888")

    # ------------------------------------------------------------------
    # Pestaña 2: Histórico diario
    # ------------------------------------------------------------------
    ws2 = wb.create_sheet("Histórico")
    ws2.sheet_view.showGridLines = False

    bank_names = sorted(set(r["bank_name"] for r in history))
    all_days   = sorted(set(r["day"] for r in history))

    # Cabecera
    ws2.cell(row=1, column=1, value="Fecha")
    _header_style(ws2.cell(row=1, column=1))
    ws2.column_dimensions["A"].width = 14

    for c, bank in enumerate(bank_names, 2):
        cell = ws2.cell(row=1, column=c, value=bank)
        _header_style(cell)
        ws2.column_dimensions[get_column_letter(c)].width = 20
    ws2.row_dimensions[1].height = 22

    # Total column
    total_col = len(bank_names) + 2
    total_header = ws2.cell(row=1, column=total_col, value="Total EUR")
    _header_style(total_header, bg="2D3748")
    ws2.column_dimensions[get_column_letter(total_col)].width = 16

    # Datos diarios
    day_map = {}
    for r in history:
        day_map.setdefault(r["day"], {})[r["bank_name"]] = r["total"]

    for r, day in enumerate(all_days, 2):
        bg = COLOR_ALT_ROW if r % 2 == 0 else "FFFFFF"
        fill = PatternFill("solid", start_color=bg)

        ws2.cell(row=r, column=1, value=day).font = Font(name="Arial", size=10)
        ws2.cell(row=r, column=1).fill = fill
        ws2.cell(row=r, column=1).alignment = Alignment(horizontal="center")

        for c, bank in enumerate(bank_names, 2):
            val = day_map.get(day, {}).get(bank, None)
            cell = ws2.cell(row=r, column=c, value=val)
            cell.fill = fill
            cell.border = _border()
            if val is not None:
                cell.number_format = '#,##0.00 €'
                cell.alignment = Alignment(horizontal="right")
            cell.font = Font(name="Arial", size=10)

        # Total row formula
        first_col = get_column_letter(2)
        last_col  = get_column_letter(len(bank_names) + 1)
        total_cell = ws2.cell(row=r, column=total_col, value=f"=SUM({first_col}{r}:{last_col}{r})")
        total_cell.number_format = '#,##0.00 €'
        total_cell.font = Font(name="Arial", bold=True, size=10)
        total_cell.fill = PatternFill("solid", start_color=COLOR_TOTAL_BG)
        total_cell.alignment = Alignment(horizontal="right")

    output = EXCEL_PATH.parent / filename
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    logger.info("Excel saved to %s", output)
    return output